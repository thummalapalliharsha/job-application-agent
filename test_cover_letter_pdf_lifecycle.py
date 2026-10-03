#!/usr/bin/env python3
"""Sandboxed regressions for the Working Cover Letter artifact lifecycle."""
from __future__ import annotations

import copy
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document

import application_assistant as aa
import career_os_api as api
import jd_resume_planner as planner

ROOT = Path(__file__).resolve().parent
APPLICATION_ID = "app_b70a2228a165"


class CoverLetterPdfLifecycleTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="cover-letter-pdf-lifecycle-")
        self.root = Path(self.temporary.name)
        self.data = self.root / "data"
        self.reports = self.root / "output" / "reports"
        self.letters = self.root / "output" / "cover_letters"
        self.resumes = self.root / "output" / "resumes"
        for directory in (self.data, self.reports, self.letters, self.resumes, self.root / "job_descriptions"):
            directory.mkdir(parents=True, exist_ok=True)
        for source in (ROOT / "data").glob("*.json"):
            shutil.copy2(source, self.data / source.name)

        original_apps = json.loads((self.data / "applications.json").read_text(encoding="utf-8"))
        target = next(item for item in original_apps["applications"] if item["application_id"] == APPLICATION_ID)
        plan_reference = target["phase8_plan_reference"].replace("\\", "/")
        source_plan = ROOT / plan_reference
        destination_plan = self.root / plan_reference
        destination_plan.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source_plan, destination_plan)

        self.final_markdown = self.letters / "existing_final.md"
        self.final_pdf = self.letters / "existing_final.pdf"
        self.final_markdown.write_text("Previously finalized cover letter.\n", encoding="utf-8")
        self.final_pdf.write_bytes(b"existing-final-pdf")
        target["cover_letter_reference"] = str(self.final_markdown.relative_to(self.root))
        target["cover_letter_pdf_reference"] = str(self.final_pdf.relative_to(self.root))
        self.resume_fields = {
            key: copy.deepcopy(value)
            for key, value in target.items()
            if key.startswith("working_resume_") or key.startswith("resume_")
        }
        (self.data / "applications.json").write_text(json.dumps(original_apps, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

        self.original_aa_paths = (aa.ROOT, aa.DATA, aa.JOBS, aa.OUT, aa.REPORTS, aa.LETTERS, aa.RESUMES)
        self.original_api_root = api.ROOT
        aa.ROOT = self.root
        aa.DATA = self.data
        aa.JOBS = self.root / "job_descriptions"
        aa.OUT = self.root / "output"
        aa.REPORTS = self.reports
        aa.LETTERS = self.letters
        aa.RESUMES = self.resumes
        api.ROOT = self.root

    def tearDown(self):
        aa.ROOT, aa.DATA, aa.JOBS, aa.OUT, aa.REPORTS, aa.LETTERS, aa.RESUMES = self.original_aa_paths
        api.ROOT = self.original_api_root
        self.temporary.cleanup()

    def test_api_generation_creates_working_markdown_docx_pdf_and_persists_refs(self):
        result = api.generate_working_cover_letter(APPLICATION_ID)
        self.assertEqual(result["decision"], "created", result)
        for reference_key in (
            "cover_letter_working_reference",
            "cover_letter_working_docx_reference",
            "cover_letter_working_pdf_reference",
        ):
            self.assertIn(reference_key, result)
            self.assertTrue((self.root / result[reference_key]).is_file())

        markdown = (self.root / result["cover_letter_working_reference"]).read_text(encoding="utf-8")
        document = Document(self.root / result["cover_letter_working_docx_reference"])
        pdf_path = self.root / result["cover_letter_working_pdf_reference"]
        pdf_bytes = pdf_path.read_bytes()
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        self.assertEqual(api._artifact_path(result["cover_letter_working_pdf_reference"]), pdf_path.resolve())
        self.assertEqual(self.final_pdf.read_bytes(), b"existing-final-pdf")
        self.assertEqual(self.final_markdown.read_text(encoding="utf-8"), "Previously finalized cover letter.\n")

        content = result["content"]
        self.assertIn("NeuralForge Technologies", content)
        self.assertIn("Junior Machine Learning Engineer", content)
        self.assertIn("applied SMOTE during preprocessing", content)
        self.assertIn("8,950-record credit-card dataset", content)
        self.assertIn("tuned the classifier with GridSearchCV", content)
        self.assertNotIn("where I worked on", content)
        self.assertNotIn("Loaded and explored an 8,950-record", content)
        self.assertNotIn("ABCD", content)
        for unsupported in ("PyTorch", "TensorFlow", "Keras", "Deep Learning", "probability expertise", "statistics expertise"):
            self.assertNotIn(unsupported.casefold(), content.casefold())
        self.assertTrue(any("Scikit-learn" in paragraph.text for paragraph in document.paragraphs))
        self.assertTrue(any("SmartFraud Classifier" in paragraph.text for paragraph in document.paragraphs))
        self.assertTrue(markdown.strip())

        store = aa.load_store()
        app = next(item for item in store["applications"] if item["application_id"] == APPLICATION_ID)
        bootstrap_app = next(item for item in api.application_payload()["applications"] if item["application_id"] == APPLICATION_ID)
        self.assertEqual(app["cover_letter_working_reference"], result["cover_letter_working_reference"])
        self.assertEqual(app["cover_letter_working_docx_reference"], result["cover_letter_working_docx_reference"])
        self.assertEqual(app["cover_letter_working_pdf_reference"], result["cover_letter_working_pdf_reference"])
        self.assertEqual(bootstrap_app["cover_letter_working_pdf_reference"], result["cover_letter_working_pdf_reference"])
        self.assertEqual(app["cover_letter_reference"], str(self.final_markdown.relative_to(self.root)))
        self.assertEqual(app["cover_letter_pdf_reference"], str(self.final_pdf.relative_to(self.root)))
        self.assertEqual({key: app.get(key) for key in self.resume_fields}, self.resume_fields)

    def test_placeholder_company_is_omitted_from_generated_letter(self):
        store = aa.load_store()
        current = next(item for item in store["applications"] if item["application_id"] == APPLICATION_ID)
        placeholder = copy.deepcopy(current)
        placeholder["application_id"] = "app_placeholder_cover_letter_test"
        placeholder["company_name"] = "ABCD"
        store["applications"].append(placeholder)
        aa.save_store(store)

        result = aa.generate_cover_letter("app_placeholder_cover_letter_test")
        self.assertEqual(result["decision"], "created")
        for placeholder_text in ("ABCD", "Company Name", "your organization", "ROLE NOT SPECIFIED"):
            self.assertNotIn(placeholder_text.casefold(), result["content"].casefold())
        self.assertIn("Junior Machine Learning Engineer", result["content"])

    def test_unmapped_project_evidence_is_composed_as_complete_prose(self):
        store = aa.load_store()
        source = next(item for item in store["applications"] if item["application_id"] == APPLICATION_ID)
        application = copy.deepcopy(source)
        application.update({
            "application_id": "app_cover_letter_prose_test",
            "company_name": "Example Analytics",
            "job_title": "Junior Data Analyst",
            "job_description_text": "Junior Data Analyst role requiring data analysis and reporting.",
            "project_selection_record_ids": ["project_imdb_movie_analysis"],
            "selected_projects": ["IMDb Movie Analysis"],
        })
        store["applications"].append(application)
        aa.save_store(store)

        result = aa.generate_cover_letter(application["application_id"])
        content = result["content"]
        self.assertIn(
            "For the IMDb Movie Analysis project, I loaded and inspected the IMDb movie dataset",
            content,
        )
        self.assertIn("I performed exploratory data analysis in a Jupyter notebook", content)
        self.assertNotIn("where I worked on", content)
        self.assertIn("My experience with", content)
        self.assertNotIn(", relevant to", content)

    def test_fresher_title_uses_jd_role_and_project_bullets_become_prose(self):
        jd = "Job Description:\nJunior Business Intelligence Analyst – Fresher\nRequired Skills:\n- Python\n- SQL\n- Pandas\nResponsibilities:\n- Build dashboards and reports."
        project_ids = [
            "project_bank_customer_clustering_dashboard",
            "project_imdb_movie_analysis",
            "project_fuel_regression_crispmlq",
        ]
        plan = planner.plan_resume(jd, aa.load_profile())
        plan["jd_analysis"]["target_role"] = "entry-level technical"
        plan_path = self.reports / "bi_role_plan.json"
        plan_path.write_text(json.dumps(plan, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

        store = aa.load_store()
        source = next(item for item in store["applications"] if item["application_id"] == APPLICATION_ID)
        application = copy.deepcopy(source)
        application.update({
            "application_id": "app_bi_role_cover_letter_test",
            "company_name": "Example Analytics",
            "job_title": "FRESHER",
            "job_description_text": jd,
            "phase8_plan_reference": "output/reports/bi_role_plan.json",
            "project_selection_record_ids": project_ids,
            "selected_projects": [
                "Bank Customer Clustering and Financial Analytics Dashboard",
                "IMDb Movie Analysis",
                "Fuel Consumption Prediction",
            ],
        })
        store["applications"].append(application)
        aa.save_store(store)

        content = aa.generate_cover_letter(application["application_id"])["content"]
        self.assertIn("I am writing to apply for the Junior Business Intelligence Analyst role", content)
        self.assertNotIn("FRESHER role", content)
        self.assertNotIn("entry-level technical role", content)
        self.assertIn("I cleaned an 8,950-record credit-card dataset", content)
        self.assertIn("I applied K-Means, Agglomerative, and DBSCAN clustering; evaluated", content)
        self.assertIn("I loaded and inspected the IMDb movie dataset", content)
        self.assertIn("I prepared flight data with numerical and categorical preprocessing", content)
        for fragment in (
            "where I worked on Loaded and explored",
            "Implemented and benchmarked K-Means",
            "where I worked on data inspection",
            "where I worked on numerical and categorical preprocessing",
        ):
            self.assertNotIn(fragment, content)
        self.assertIn(
            "My experience with Python, Pandas, and SQL is relevant to data analysis and SQL-based workflows.",
            content,
        )

    def test_cover_letter_does_not_mention_unsupported_sql(self):
        jd = "Job Description:\nJunior Business Intelligence Analyst\nRequired Skills:\n- Python\nResponsibilities:\n- Build dashboards and reports."
        plan = planner.plan_resume(jd, aa.load_profile())
        plan_path = self.reports / "bi_no_sql_plan.json"
        plan_path.write_text(json.dumps(plan, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        store = aa.load_store()
        source = next(item for item in store["applications"] if item["application_id"] == APPLICATION_ID)
        application = copy.deepcopy(source)
        application.update({
            "application_id": "app_bi_no_sql_cover_letter_test",
            "job_title": "FRESHER",
            "job_description_text": jd,
            "phase8_plan_reference": "output/reports/bi_no_sql_plan.json",
            "project_selection_record_ids": ["project_bank_customer_clustering_dashboard"],
            "selected_projects": ["Bank Customer Clustering and Financial Analytics Dashboard"],
        })
        store["applications"].append(application)
        aa.save_store(store)

        content = aa.generate_cover_letter(application["application_id"])["content"]
        self.assertNotIn("SQL", content)

    def test_react_consumes_api_pdf_reference_and_renders_working_download(self):
        source = (ROOT / "frontend" / "src" / "main.tsx").read_text(encoding="utf-8")
        self.assertIn("setWorkingPdfRef(result.cover_letter_working_pdf_reference)", source)
        self.assertIn("{workingPdfRef && (", source)
        self.assertIn("encodeURIComponent(workingPdfRef)", source)
        self.assertIn("DOWNLOAD WORKING PDF", source)
        self.assertIn("download>", source)


if __name__ == "__main__":
    unittest.main(verbosity=2)
