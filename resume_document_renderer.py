"""Additive, editor-specific DOCX/PDF renderer for validated ResumeDocument v1.

The module reuses the frozen resume generator's ATS template, body-clearing,
font conventions, heading helper, hyperlink helper, and LibreOffice/Poppler
conversion. It never writes to the frozen generator's default Working/Final
paths; callers must supply isolated staging destinations.
"""
from __future__ import annotations

import shutil
import subprocess
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.text.run import Run
from docx.oxml.ns import qn

import resume_document_model as rdm
import resume_generator as rg

_ALIGNMENT = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_DISALLOWED_DOCX_TAGS = {
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl": "tables",
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent": "text boxes",
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing": "drawings/images",
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pict": "legacy images/shapes",
    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}object": "embedded objects",
    "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor": "floating objects",
    "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline": "inline objects",
}


def validate_docx_structure(docx_path: str | Path) -> dict[str, bool]:
    """Fail closed on non-ATS tables, shapes, text boxes, images, or objects."""
    path = Path(docx_path)
    found: set[str] = set()
    story_prefixes = ("word/document.xml", "word/header", "word/footer", "word/footnotes.xml", "word/endnotes.xml", "word/comments.xml")
    with zipfile.ZipFile(path) as package:
        for name in package.namelist():
            if not name.startswith(story_prefixes) or not name.endswith(".xml"):
                continue
            root = ET.fromstring(package.read(name))
            for element in root.iter():
                category = _DISALLOWED_DOCX_TAGS.get(element.tag)
                if category:
                    found.add(category)
    if found:
        raise ValueError("Rendered DOCX contains unsupported ATS structure: " + ", ".join(sorted(found)))
    return {"single_column": True, "no_tables": True, "no_text_boxes": True,
            "no_drawings_images_or_floating_objects": True}


def _apply_paragraph_format(paragraph: Any, formatting: dict[str, Any]) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(float(formatting.get("space_before_pt", 0)))
    pf.space_after = Pt(float(formatting.get("space_after_pt", 0)))
    pf.line_spacing = float(formatting.get("line_spacing", 1.0))
    if "alignment" in formatting:
        paragraph.alignment = _ALIGNMENT[formatting["alignment"]]
    if "left_indent_in" in formatting:
        pf.left_indent = Inches(float(formatting["left_indent_in"]))
    if "first_line_indent_in" in formatting:
        pf.first_line_indent = Inches(float(formatting["first_line_indent_in"]))


def _apply_run_format(run: Any, size: float, formatting: dict[str, Any], marks: set[str], *, hyperlink: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(float(size))
    run.bold = bool(formatting.get("bold", False) or "bold" in marks)
    run.italic = bool(formatting.get("italic", False) or "italic" in marks)
    run.underline = bool(hyperlink or formatting.get("underline", False) or "underline" in marks)


def _append_runs(paragraph: Any, runs: list[dict[str, Any]], formatting: dict[str, Any]) -> None:
    size = float(formatting.get("font_size_pt", 11.0))
    for item in runs:
        text = str(item.get("text", ""))
        marks = {str(mark.get("type")) for mark in item.get("marks", []) if isinstance(mark, dict)}
        link = next((mark.get("href") for mark in item.get("marks", [])
                     if isinstance(mark, dict) and mark.get("type") == "link"), None)
        if link:
            before = len(paragraph._p)
            rg.add_hyperlink(paragraph, text, str(link), size=size,
                             bold=bool(formatting.get("bold", False) or "bold" in marks))
            # The frozen helper creates one hyperlink run; apply only the other
            # supported run attributes while preserving its blue/underlined link.
            for element in list(paragraph._p)[before:]:
                if element.tag.endswith("}hyperlink"):
                    for run_element in element.xpath(".//w:r"):
                        _apply_run_format(Run(run_element, paragraph), size, formatting, marks, hyperlink=True)
        else:
            run = paragraph.add_run(text)
            _apply_run_format(run, size, formatting, marks)


def _new_paragraph(document: Any, formatting: dict[str, Any]) -> Any:
    paragraph = document.add_paragraph()
    _apply_paragraph_format(paragraph, formatting)
    return paragraph


def _render_block(document: Any, block: dict[str, Any], bullet_counters: dict[str, int]) -> None:
    kind = block["type"]
    formatting = block.get("formatting", {})

    if kind == "layout_paragraph":
        _new_paragraph(document, formatting)
        return

    if kind in {"skill_group", "project_tech_stack"}:
        paragraph = _new_paragraph(document, formatting)
        label_formatting = block.get("label_formatting", {})
        label = str(block.get("label", ""))
        if kind == "skill_group":
            label += ": "
        label_run = paragraph.add_run(label)
        _apply_run_format(label_run, float(label_formatting.get("font_size_pt", 12.0)),
                          label_formatting, set())
        values = ", ".join(str(item["text"]) for item in block.get("items", []))
        value_run = paragraph.add_run(values)
        _apply_run_format(value_run, float(formatting.get("font_size_pt", 11.0)), formatting, set())
        return

    paragraph = _new_paragraph(document, formatting)
    if kind in {"project_bullet", "experience_bullet", "certification_entry"}:
        style = formatting.get("list_style", "bullet")
        if style == "bullet":
            marker = "• "
        elif style == "ordered":
            group = str(block.get("project_id") or block.get("experience_id") or "certification")
            bullet_counters[group] = bullet_counters.get(group, 0) + 1
            marker = f"{bullet_counters[group]}. "
        else:
            marker = ""
        marker_run = paragraph.add_run(marker)
        _apply_run_format(marker_run, float(formatting.get("font_size_pt", 11.0)), formatting, set())
    if "runs" in block:
        _append_runs(paragraph, block["runs"], formatting)
    else:
        raise ValueError(f"Unsupported ResumeDocument block without text runs: {kind}")


def render_docx(document: dict[str, Any], output_path: str | Path) -> Path:
    """Render the already-validated structured model to a new DOCX path."""
    output = Path(output_path)
    if output.suffix.lower() != ".docx":
        raise ValueError("DOCX output path must use the .docx extension")
    output.parent.mkdir(parents=True, exist_ok=True)

    # No blank-document fallback: the approved ATS template is a hard requirement.
    doc = Document(str(rg.TEMPLATE))
    rg.clear_template_body(doc)
    for section in document["content"]["sections"]:
        title = section.get("title")
        if title:
            # Keep the frozen heading/rule helper, then apply only bounded model formatting.
            rg.section(doc, str(title))
            heading = doc.paragraphs[-1]
            heading_formatting = section.get("formatting", {})
            _apply_paragraph_format(heading, heading_formatting)
            for run in heading.runs:
                _apply_run_format(run, float(heading_formatting.get("font_size_pt", 13.0)),
                                  heading_formatting, set())
        counters: dict[str, int] = {}
        for block in section["blocks"]:
            _render_block(doc, block, counters)

    doc.save(str(output))
    # Re-open the file so corrupt/incomplete packages fail before PDF conversion.
    reopened = Document(str(output))
    if reopened.tables:
        raise ValueError("Rendered DOCX unexpectedly contains a table")
    validate_docx_structure(output)
    return output


def render_pdf(docx_path: str | Path, pdf_path: str | Path, workdir: str | Path) -> dict[str, Any]:
    """Convert the staged DOCX with the frozen generator's LibreOffice/Poppler path."""
    docx = Path(docx_path)
    pdf = Path(pdf_path)
    work = Path(workdir)
    if docx.suffix.lower() != ".docx" or pdf.suffix.lower() != ".pdf":
        raise ValueError("Renderer requires .docx and .pdf output paths")
    work.mkdir(parents=True, exist_ok=True)
    pdf.parent.mkdir(parents=True, exist_ok=True)
    page_count, _generator_text, converted = rg.render_page_count(docx, work)
    if converted.resolve() != pdf.resolve():
        shutil.copy2(converted, pdf)
    if not pdf.is_file() or pdf.stat().st_size == 0:
        raise ValueError("LibreOffice did not produce a non-empty PDF")
    # Poppler emits UTF-8, but subprocess.run(text=True) in the frozen helper
    # inherits Windows' legacy code page. Decode explicitly so Unicode list
    # markers and punctuation compare faithfully in the save-time text check.
    extraction = subprocess.run([rg.resolve_executable("pdftotext"), str(pdf), "-"],
                                check=True, capture_output=True, encoding="utf-8")
    text = extraction.stdout
    return {"pdf_path": pdf, "page_count": page_count, "extracted_text": text}


def render_resume_document(document: dict[str, Any], application: dict[str, Any], plan: dict[str, Any],
                           profile: dict[str, Any], docx_path: str | Path, pdf_path: str | Path,
                           workdir: str | Path) -> dict[str, Any]:
    """Validate model schema/context, then render DOCX and PDF to staging paths."""
    model_report = rdm.validate_resume_document(document, application, plan, profile)
    docx = render_docx(document, docx_path)
    structure_report = validate_docx_structure(docx)
    pdf_report = render_pdf(docx, pdf_path, workdir)
    return {"model_validation": model_report, "docx_path": docx,
            "pdf_path": pdf_report["pdf_path"], "page_count": pdf_report["page_count"],
            "extracted_text": pdf_report["extracted_text"], "structure_validation": structure_report}
