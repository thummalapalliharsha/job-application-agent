from __future__ import annotations

import hashlib
import json
import copy
import hmac
import mimetypes
import os
import re
import shutil
import threading
import tempfile
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote_to_bytes, urlparse

import application_assistant as aa
import jd_resume_planner as planner
import profile_update_agent as pua
from career_os_config import (
    DATA_DIR,
    FRONTEND_DIST_DIR,
    OUTPUT_DIR,
    STORAGE_ROOT,
    production_storage_is_configured,
    resolve_storage_reference,
    storage_reference,
    storage_root,
)
try:
    import resume_generator as rg
except ModuleNotFoundError:
    rg = None

ROOT = STORAGE_ROOT
FRONTEND = FRONTEND_DIST_DIR
_RESUME_DOCUMENT_LOCK = threading.RLock()
_LOCAL_ALLOWED_HOSTS = {
    f"{host}:{port}"
    for host in ("localhost", "127.0.0.1")
    for port in (8504, 5173, 4173)
}
_LOCAL_ALLOWED_ORIGINS = {f"http://{host}" for host in _LOCAL_ALLOWED_HOSTS}
_API_TOKEN = os.environ.get("CAREER_OS_API_TOKEN", "").strip()
_RUNTIME_ENV = os.environ.get("CAREER_OS_ENV", "development").strip().casefold()
_BAD_PERCENT_ESCAPE = re.compile(r"%(?![0-9A-Fa-f]{2})")
_WINDOWS_DRIVE = re.compile(r"^[A-Za-z]:")


def _environment_items(name: str) -> set[str]:
    return {item.strip() for item in os.environ.get(name, "").split(",") if item.strip()}


def _configured_origins() -> set[str]:
    origins = set(_LOCAL_ALLOWED_ORIGINS)
    for origin in _environment_items("CAREER_OS_ALLOWED_ORIGINS"):
        parsed = urlparse(origin)
        if (origin != "*" and parsed.scheme in {"http", "https"} and parsed.netloc
                and parsed.username is None and parsed.password is None
                and parsed.path in {"", "/"} and not parsed.params and not parsed.query and not parsed.fragment):
            origins.add(f"{parsed.scheme}://{parsed.netloc}")
    return origins


_ALLOWED_ORIGINS = _configured_origins()
_ALLOWED_HOSTS = _LOCAL_ALLOWED_HOSTS | _environment_items("CAREER_OS_ALLOWED_HOSTS")
_ALLOWED_HOSTS.update(urlparse(origin).netloc for origin in _ALLOWED_ORIGINS)

def _api_server_address() -> tuple[str, int]:
    production = _RUNTIME_ENV in {"production", "prod"}
    host = os.environ.get(
        "CAREER_OS_HOST",
        "0.0.0.0" if production else "127.0.0.1",
    ).strip()
    port = int(os.environ.get("PORT", os.environ.get("CAREER_OS_PORT", "8504")))
    if not host or not 1 <= port <= 65535:
        raise ValueError("CAREER_OS_HOST and PORT must specify a valid bind address")
    return host, port


def _valid_api_token(method: str, path: str, authorization: str | None) -> bool:
    if not _API_TOKEN or method == "OPTIONS" or path == "/api/health":
        return True
    if not isinstance(authorization, str) or not authorization.startswith("Bearer "):
        return False
    return hmac.compare_digest(authorization[7:], _API_TOKEN)


def _decode_url_path(value):
    """Strictly decode an origin-form request path, or return None."""
    if not isinstance(value, str) or not value.startswith("/") or _BAD_PERCENT_ESCAPE.search(value):
        return None
    try:
        decoded = unquote_to_bytes(value).decode("utf-8", "strict")
    except (UnicodeDecodeError, ValueError):
        return None
    if any(ord(char) < 32 or ord(char) == 127 for char in decoded) or "\\" in decoded or decoded.startswith("//"):
        return None
    relative = decoded[1:]
    if relative.startswith("/") or _WINDOWS_DRIVE.match(relative):
        return None
    if any(part in {".", ".."} for part in relative.split("/")):
        return None
    return decoded


def _contained_file(base, relative):
    """Resolve a relative path while containing symlink targets in base."""
    try:
        resolved_base = base.resolve()
        candidate = (base / relative).resolve()
        candidate.relative_to(resolved_base)
    except (OSError, RuntimeError, ValueError):
        return None
    return candidate if candidate.is_file() else None


def _storage_root():
    return storage_root(ROOT)


def _resolve_storage_reference(reference):
    return resolve_storage_reference(reference, ROOT)


def _storage_reference(path):
    resolved = Path(path).resolve()
    try:
        return resolved.relative_to(_storage_root()).as_posix()
    except ValueError:
        return storage_reference(resolved)


def _artifact_path(reference):
    """Resolve a stored artifact reference into one of the public artifact trees."""
    if (not isinstance(reference, str) or not reference
            or any(ord(char) < 32 or ord(char) == 127 for char in reference)):
        return None
    if reference.startswith(("/", "\\")) or reference.startswith("//") or reference.startswith("\\\\"):
        return None
    normalized = reference.replace("\\", "/")
    if normalized.startswith("/") or _WINDOWS_DRIVE.match(normalized):
        return None
    parts = normalized.split("/")
    if any(part in {"", ".", ".."} for part in parts):
        return None

    allowed = (
        (Path("output") / "resumes", {".docx", ".pdf"}),
        (Path("output") / "cover_letters", {".md", ".docx", ".pdf"}),
    )
    for prefix, extensions in allowed:
        try:
            relative = Path(normalized).relative_to(prefix)
        except ValueError:
            continue
        if relative.suffix.lower() not in extensions:
            return None
        return _contained_file(_storage_root() / prefix, relative)
    return None


def profile_summary():
    profile = planner.load_profile()
    master = profile.get("master_profile", {}).get("profile", {})
    projects = profile.get("projects", {}).get("projects", [])
    skills = [s for group in profile.get("skills", {}).get("skill_groups", []) for s in group.get("skills", [])]
    return {"name": master.get("name"), "headline": master.get("headline") or master.get("summary"), "location": master.get("location"), "projects": projects, "skills": skills, "education": profile.get("education", {}).get("education", []), "experience": profile.get("experience", {}).get("experiences", []), "certifications": profile.get("certifications", {}).get("certifications", [])}


def confirm_skill_gap(aid, payload):
    """Persist an explicitly confirmed skill as candidate-provided and stale the application plan."""
    if not isinstance(payload, dict) or payload.get("confirmed") is not True:
        return {"decision": "confirmation_required", "message": "Explicit confirmation is required before adding a candidate-provided skill."}
    skill_name = str(payload.get("skill") or "").strip()
    category = str(payload.get("category") or "").strip()
    if not skill_name or len(skill_name) > 100:
        return {"decision": "invalid", "message": "A valid skill name is required."}

    with _RESUME_DOCUMENT_LOCK:
        app = get_application(aid)
        if not app:
            return {"decision": "not_found"}
        plan_path = _safe_project_path(app.get("phase8_plan_reference"))
        if not plan_path or not plan_path.is_file():
            return {"decision": "error", "message": "The current Resume Plan could not be safely resolved."}
        try:
            plan = json.loads(plan_path.read_text(encoding="utf-8"))
            matching = next((item for item in plan.get("candidate_matching", [])
                             if planner.norm(item.get("requirement")) == planner.norm(skill_name)), None)
            if not matching or matching.get("classification") != "required" or matching.get("evidence_status") not in {"UNSUPPORTED", "UNKNOWN"}:
                return {"decision": "invalid", "message": "Only a currently unsupported required JD skill can be added through this confirmation."}
            store_path = aa.DATA / "applications.json"
            original_store = store_path.read_bytes()
            original_plan = plan_path.read_bytes()
            profile_paths = [_storage_root() / "data" / "skills.json", _storage_root() / "data" / "master_profile.json"]
            original_profile = {path: path.read_bytes() for path in profile_paths}
            profile = pua.all_data(ROOT)
            update_plan = pua.plan_skill_gap_addition(skill_name, category, root=ROOT, application_id=aid)
            if update_plan.get("decision") != "planned":
                return {"decision": update_plan.get("decision", "invalid"), "message": "The skill could not be planned for profile update.", "details": update_plan}
            pua.apply_plan(update_plan, root=ROOT, confirm=True)
            profile = pua.all_data(ROOT)
            updated_plan = planner.plan_resume(app.get("job_description_text", ""), profile)
            automatic = copy.deepcopy(updated_plan["resume_plan"].get("projects_to_include", []))
            updated_plan["resume_plan"]["automatic_projects_to_include"] = copy.deepcopy(automatic)
            mode = app.get("project_selection_mode") or "automatic"
            selected = automatic
            if mode == "manual":
                eligible = {item.get("record_id"): item for item in profile["projects"].get("projects", [])
                            if item.get("status") == "verified" and item.get("project_status") == "completed"}
                selected = [eligible[record_id] for record_id in app.get("project_selection_record_ids", []) if record_id in eligible]
                if selected:
                    updated_plan["resume_plan"]["projects_to_include"] = selected
                    updated_plan["resume_plan"]["project_selection_source"] = "manual"
                    updated_plan["resume_plan"]["project_selection_record_ids"] = [item["record_id"] for item in selected]
                else:
                    mode = "automatic"
            updated_plan.setdefault("approval_checkpoint", {})["resume_generation_allowed"] = False
            _atomic_write_json(plan_path, updated_plan)

            store = json.loads(store_path.read_bytes().decode("utf-8"))
            target = next(item for item in store.get("applications", []) if item.get("application_id") == aid)
            selected = updated_plan["resume_plan"].get("projects_to_include", automatic)
            target.update({
                "selected_projects": [item.get("name") for item in selected],
                "project_selection_mode": mode,
                "project_selection_source": mode,
                "project_selection_record_ids": [item.get("record_id") for item in selected],
                "selected_skills": [item.get("name") for item in updated_plan["resume_plan"].get("skills_to_include", [])],
                "selected_certifications": [item.get("name") for item in updated_plan["resume_plan"].get("certifications_to_include", [])],
                "experience_decision": updated_plan["resume_plan"].get("experience_decisions", []),
                "requirements_summary": updated_plan["jd_analysis"].get("requirements", []),
                "supported_requirements": [item.get("requirement") for item in updated_plan["evidence_summary"].get("supported_requirements", [])],
                "partial_requirements": [item.get("requirement") for item in updated_plan["evidence_summary"].get("partial_requirements", [])],
                "unsupported_requirements": [item.get("requirement") for item in updated_plan["evidence_summary"].get("unsupported_requirements", [])],
                "candidate_gap_summary": aa.gap_summary(updated_plan),
                "application_checklist": aa.checklist(updated_plan),
                "resume_generation_allowed": False,
                "resume_working_artifact_stale": bool(target.get("working_resume_generation_id") or target.get("working_resume_docx_path")),
                "current_status": "awaiting_resume_approval",
                "last_updated": aa.now(),
            })
            _write_application_store_atomically(original_store, store)
            return {"decision": "skill_added_candidate_provided", "message": "Added as candidate-provided evidence. It will not appear in a resume until verified. The refreshed Resume Plan requires approval before regeneration.", "skill": skill_name, "status": "candidate_provided", "application": target, "plan": updated_plan}
        except Exception as exc:
            for path, content in original_profile.items() if 'original_profile' in locals() else []:
                path.write_bytes(content)
            if 'original_plan' in locals():
                plan_path.write_bytes(original_plan)
            return {"decision": "error", "message": f"Skill-gap confirmation was not completed: {type(exc).__name__}: {exc}"}


def application_payload():
    apps = aa.load_store().get("applications", [])
    counts = {}
    for item in apps:
        status = item.get("current_status", "unknown")
        counts[status] = counts.get(status, 0) + 1
    enriched = []
    for item in apps:
        record = dict(item)
        for prefix, doc_key, pdf_key in (
            ("working_resume", "working_resume_docx_path", "working_resume_pdf_path"),
            ("final_resume", "resume_docx_path", "resume_pdf_path"),
        ):
            doc = resolve_ref(record.get(doc_key) or record.get("working_resume_reference" if prefix == "working_resume" else "resume_reference")) if (record.get(doc_key) or record.get("working_resume_reference" if prefix == "working_resume" else "resume_reference")) else None
            pdf = resolve_ref(record.get(pdf_key) or record.get("working_resume_pdf_reference" if prefix == "working_resume" else "resume_pdf_reference")) if (record.get(pdf_key) or record.get("working_resume_pdf_reference" if prefix == "working_resume" else "resume_pdf_reference")) else None
            record[f"{prefix}_docx_available"] = bool(doc and doc.is_file())
            record[f"{prefix}_pdf_available"] = bool(pdf and pdf.is_file())
        enriched.append(record)
    return {"applications": enriched, "counts": counts}


def get_application(aid):
    return next((item for item in aa.load_store().get("applications", []) if item.get("application_id") == aid), None)


def _safe_project_path(reference):
    if not isinstance(reference, str) or not reference.strip():
        return None
    return _resolve_storage_reference(reference)


def _resume_document_context(app):
    """Return the fresh approved baseline and, if active/current, its saved revision."""
    if rg is None:
        return {"decision": "backend_unavailable", "error": "Resume document tooling is unavailable."}
    import resume_document_model as rdm
    import resume_document_validation as rdv
    try:
        baseline = rdm.build_resume_document_for_application(app, root=ROOT)
        plan_path = _safe_project_path(baseline["source_plan_reference"])
        if not plan_path or not plan_path.is_file():
            return {"decision": "invalid", "error": "The approved Resume Plan could not be safely resolved."}
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        profile = rg.profile()
        document = baseline
        active_doc_ref = app.get("working_resume_docx_path") or app.get("working_resume_reference")
        active_pdf_ref = app.get("working_resume_pdf_path") or app.get("working_resume_pdf_reference")
        sidecar_ref = app.get("working_resume_revision_reference")
        sidecar_path = _safe_project_path(sidecar_ref) if sidecar_ref else None
        sidecar = None
        if sidecar_path and sidecar_path.is_file():
            candidate_sidecar = json.loads(sidecar_path.read_text(encoding="utf-8"))
            # A prior edited sidecar can remain after a deliberate regeneration.
            # It is active only while its DOCX path is the app's current Working path.
            if active_doc_ref and active_doc_ref == candidate_sidecar.get("working_docx_reference"):
                sidecar = candidate_sidecar
                if app.get("resume_working_artifact_stale"):
                    return {"decision": "stale_revision", "error": "This edited Working Resume is stale. Approve the current Resume Plan and regenerate before editing or finalizing."}
                if sidecar.get("application_id") != app.get("application_id") or sidecar.get("revision_id") != app.get("working_resume_revision_id"):
                    return {"decision": "stale_revision", "error": "The saved revision metadata does not match the selected application. Regenerate the Working Resume before continuing."}
                saved_document = sidecar.get("document")
                if not isinstance(saved_document, dict) or saved_document.get("revision_id") != sidecar.get("revision_id"):
                    return {"decision": "invalid", "error": "The saved ResumeDocument sidecar is incomplete or inconsistent."}
                if saved_document.get("source_plan_sha256") != rdm._plan_digest(plan):
                    return {"decision": "stale_revision", "error": "The approved Resume Plan changed after this Working Resume revision was saved."}
                if saved_document.get("source_generation_reference") != app.get("working_resume_generation_id"):
                    return {"decision": "stale_revision", "error": "The saved Working Resume no longer matches the current generation."}
                docx_path = _safe_project_path(active_doc_ref)
                pdf_path = _safe_project_path(active_pdf_ref)
                if not docx_path or not pdf_path or not docx_path.is_file() or not pdf_path.is_file():
                    return {"decision": "artifact_missing", "error": "A file for the saved Working Resume revision is missing."}
                if file_sha256(docx_path) != sidecar.get("working_docx_sha256") or file_sha256(pdf_path) != sidecar.get("working_pdf_sha256"):
                    return {"decision": "artifact_modified", "error": "A saved Working Resume file changed after validation; restore or regenerate it before editing."}
                if app.get("working_resume_docx_sha256") != sidecar.get("working_docx_sha256") or app.get("working_resume_pdf_sha256") != sidecar.get("working_pdf_sha256"):
                    return {"decision": "artifact_modified", "error": "The active Working Resume hash metadata does not match its saved revision."}
                validation_record = sidecar.get("validation")
                if not isinstance(validation_record, dict) or validation_record.get("final_status") != "PASS":
                    return {"decision": "invalid", "error": "The saved revision has no passing validation record."}
                rdm.validate_resume_document(saved_document, app, plan, profile)
                policy = rdv.validate_edit_policy(saved_document, baseline, profile)
                if not policy.get("valid"):
                    return {"decision": "invalid", "error": "The saved revision no longer satisfies source-evidence checks.", "validation": policy}
                document = saved_document
        return {"decision": "ready", "application": app, "document": document,
                "baseline_document": baseline, "plan": plan, "profile": profile, "sidecar": sidecar}
    except rdm.ResumeDocumentValidationError as exc:
        return {"decision": "invalid", "error": str(exc)}
    except (KeyError, TypeError, ValueError, OSError) as exc:
        return {"decision": "invalid", "error": f"ResumeDocument context could not be loaded: {exc}"}


def resume_document_payload(aid):
    """Load the selected application's current source-backed ResumeDocument."""
    app = get_application(aid)
    if not app:
        return {"decision": "not_found"}
    context = _resume_document_context(app)
    if context.get("decision") != "ready":
        return context
    import resume_document_model as rdm
    projects, _, _ = rdm._approved_context(app, context["plan"], context["profile"])
    skill_templates = [copy.deepcopy(block) for section in context["baseline_document"]["content"]["sections"]
                       if section.get("type") == "skills" for block in section["blocks"]
                       if block.get("type") == "skill_group"]
    return {"decision": "ready", "application_id": aid, "company_name": app.get("company_name"),
            "document": context["document"], "base_revision_id": context["document"].get("revision_id"),
            "persisted": bool(context.get("sidecar")),
            "skill_catalog": rdm.canonical_skill_catalog(context["profile"], projects),
            "skill_group_templates": skill_templates}


def validate_resume_document_payload(aid, payload):
    """Validate an editor revision in memory without persisting it."""
    app = get_application(aid)
    if not app:
        return {"decision": "not_found"}
    if rg is None:
        return {"decision": "backend_unavailable", "valid": False,
                "error": "Resume document tooling is unavailable in this runtime."}
    if not isinstance(payload, dict) or not isinstance(payload.get("document"), dict):
        return {"decision": "invalid", "valid": False, "error": "A ResumeDocument object is required."}
    import resume_document_model as rdm
    import resume_document_validation as rdv
    try:
        context = _resume_document_context(app)
        if context.get("decision") != "ready":
            return {"decision": context.get("decision", "invalid"), "valid": False,
                    "error": context.get("error", "The current ResumeDocument context is unavailable."), "persisted": False}
        baseline = context["document"]
        model_baseline = context["baseline_document"]
        plan, profile = context["plan"], context["profile"]
        edit_context = payload.get("edit_context") if isinstance(payload.get("edit_context"), dict) else {}
        supplied_base = payload.get("base_revision_id") or edit_context.get("original_revision_id")
        if supplied_base and supplied_base != baseline.get("revision_id"):
            return {"decision": "stale_revision", "valid": False,
                    "error": "This edit is based on an older ResumeDocument revision. Reload the current Working Resume before saving.",
                    "persisted": False}
        candidate = copy.deepcopy(payload["document"])
        candidate["revision_id"] = rdm._revision_id(candidate)
        report = rdm.validate_resume_document(candidate, app, plan, profile)
        policy = rdv.validate_edit_policy(candidate, model_baseline, profile)
        if not policy.get("valid"):
            return {"decision": "invalid", "valid": False, "error": "Source-evidence validation failed.",
                    "validation": report, "claim_validation": policy, "persisted": False}
    except rdm.ResumeDocumentValidationError as exc:
        return {"decision": "invalid", "valid": False, "error": str(exc)}
    except (KeyError, TypeError, ValueError) as exc:
        return {"decision": "invalid", "valid": False, "error": f"Invalid ResumeDocument: {exc}"}
    return {"decision": "valid", "valid": True, "document": candidate,
            "validation": report, "claim_validation": policy, "persisted": False}


def resolve_ref(reference):
    return _resolve_storage_reference(reference)


def file_sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _atomic_write_json(path, value, *, exclusive=False):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if exclusive and path.exists():
        raise FileExistsError(str(path))
    fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(json.dumps(value, indent=2, ensure_ascii=False) + "\n")
            handle.flush()
            os.fsync(handle.fileno())
        if exclusive and path.exists():
            raise FileExistsError(str(path))
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _write_application_store_atomically(expected_bytes, store):
    store_path = aa.DATA / "applications.json"
    if not store_path.is_file() or store_path.read_bytes() != expected_bytes:
        raise RuntimeError("Application data changed during validation. Reload the current application and retry.")
    fd, temporary = tempfile.mkstemp(prefix=".applications.resume-save.", suffix=".tmp", dir=str(store_path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            handle.write(json.dumps(store, indent=2, ensure_ascii=False) + "\n")
            handle.flush()
            os.fsync(handle.fileno())
        if store_path.read_bytes() != expected_bytes:
            raise RuntimeError("Application data changed during validation. Reload the current application and retry.")
        os.replace(temporary, store_path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _normalized_document_text(text):
    import re
    return re.sub(r"\s+", "", str(text or "")).casefold()


def _build_resume_change_ledger(baseline, candidate):
    def index(document):
        return {block["id"]: block for section in document["content"]["sections"] for block in section["blocks"]}
    before, after = index(baseline), index(candidate)
    entries = []
    for block_id, new in after.items():
        old = before.get(block_id)
        if old is None or new.get("edit_policy") != "editable_source_backed":
            continue
        old_text = "".join(run.get("text", "") for run in old.get("runs", []))
        new_text = "".join(run.get("text", "") for run in new.get("runs", []))
        formatting_changed = old.get("formatting") != new.get("formatting") or [run.get("marks") for run in old.get("runs", [])] != [run.get("marks") for run in new.get("runs", [])]
        if old_text != new_text or formatting_changed:
            entries.append({"block_id": block_id, "block_type": new.get("type"), "original_text": old_text,
                            "edited_text": new_text, "source_refs": copy.deepcopy(new.get("source_refs", [])),
                            "formatting_changed": formatting_changed})
    return entries


def _save_resume_response(application, document, report, docx_ref, pdf_ref, saved_at, *, idempotent=False):
    return {
        "decision": "saved", "persisted": True, "idempotent": idempotent,
        "message": "Working Resume saved and validated.", "application": application,
        "working": {"docx_reference": docx_ref, "pdf_reference": pdf_ref,
                    "revision_id": document["revision_id"], "saved_at": saved_at,
                    "page_count": report.get("page_count", 1)},
        "validation": {"final_status": report.get("final_status"),
                       "page_count": report.get("page_count"),
                       "checks": report.get("checks", {}),
                       "claim_truth_semantically_verified": False},
        "previous_working_preserved": True,
    }


def _frozen_generator_validation_view(document, baseline_document):
    """Normalize presentation-only labels for the frozen generator's legacy checks.

    The final candidate remains the rendered/activated artifact. The frozen
    validator predates editable headings, project display titles, and language
    skill text, so its temporary view restores only those labels while retaining
    the candidate's structure, factual text, and layout settings.
    """
    import resume_document_model as rdm
    view = copy.deepcopy(document)
    baseline_sections = {section["id"]: section for section in baseline_document["content"]["sections"]}
    baseline_blocks = {block["id"]: block for section in baseline_document["content"]["sections"] for block in section["blocks"]}
    for section in view["content"]["sections"]:
        base_section = baseline_sections[section["id"]]
        section["title"] = base_section["title"]
        for block in section["blocks"]:
            base_block = baseline_blocks.get(block["id"])
            if block.get("type") == "project_entry" and base_block:
                block["runs"] = copy.deepcopy(base_block["runs"])
            elif section.get("type") == "skills" and block.get("type") == "skill_group" and block.get("label") == "Languages":
                base_group = baseline_blocks.get(block["id"])
                if base_group:
                    block["items"] = copy.deepcopy(base_group["items"])
                    block["source_refs"] = copy.deepcopy(base_group["source_refs"])
    view["source_references"] = rdm._collect_refs(view["content"])
    view["revision_id"] = rdm._revision_id(view)
    return view


def save_resume_document_payload(aid, payload):
    """Validate, render, and atomically activate a new Working Resume revision."""
    if rg is None:
        return {"decision": "backend_unavailable", "message": "Resume document tooling is unavailable."}
    if not isinstance(payload, dict) or not isinstance(payload.get("document"), dict):
        return {"decision": "invalid", "message": "A ResumeDocument object is required.", "previous_working_preserved": True}

    import resume_document_model as rdm
    import resume_document_renderer as rdr
    import resume_document_validation as rdv

    with _RESUME_DOCUMENT_LOCK:
        store_path = aa.DATA / "applications.json"
        try:
            original_store_bytes = store_path.read_bytes()
            store = json.loads(original_store_bytes.decode("utf-8"))
            app = next((item for item in store.get("applications", []) if item.get("application_id") == aid), None)
        except (OSError, json.JSONDecodeError) as exc:
            return {"decision": "error", "message": f"Application data could not be read: {exc}", "previous_working_preserved": True}
        if not app:
            return {"decision": "not_found"}
        if not app.get("resume_generation_allowed"):
            return {"decision": "approval_required", "message": "Approve the current Resume Plan before saving a Working Resume.", "previous_working_preserved": True}
        if app.get("resume_working_artifact_stale"):
            return {"decision": "stale_revision", "message": "The selected application’s plan or project selection changed. Approve the current plan and regenerate before editing.", "previous_working_preserved": True}

        context = _resume_document_context(app)
        if context.get("decision") != "ready":
            return {"decision": context.get("decision", "invalid"), "message": context.get("error", "The current ResumeDocument context is unavailable."), "previous_working_preserved": True}
        current_document = context["document"]
        baseline_document = context["baseline_document"]
        supplied_base = payload.get("base_revision_id")
        if supplied_base != current_document.get("revision_id"):
            return {"decision": "stale_revision", "message": "This edit is based on an older ResumeDocument revision. Reload the current Working Resume before saving.", "previous_working_preserved": True}

        candidate = copy.deepcopy(payload["document"])
        candidate["status"] = "working"
        validation_app = copy.deepcopy(app)
        new_generation = not validation_app.get("working_resume_generation_id")
        if new_generation:
            validation_app["working_resume_generation_id"] = "gen_" + uuid.uuid4().hex
            candidate["source_generation_reference"] = validation_app["working_resume_generation_id"]
        candidate["revision_id"] = rdm._revision_id(candidate)
        try:
            model_report = rdm.validate_resume_document(candidate, validation_app, context["plan"], context["profile"])
            policy_report = rdv.validate_edit_policy(candidate, baseline_document, context["profile"])
        except rdm.ResumeDocumentValidationError as exc:
            return {"decision": "invalid", "message": str(exc), "previous_working_preserved": True}
        if not policy_report.get("valid"):
            return {"decision": "validation_failed", "message": "Source-evidence validation failed. Remove unsupported additions and retry.",
                    "errors": policy_report.get("errors", []), "claim_validation": policy_report,
                    "previous_working_preserved": True}

        # Saving an unchanged already-active revision is idempotent and creates no files.
        current_sidecar = context.get("sidecar")
        if current_sidecar and candidate["revision_id"] == current_document.get("revision_id"):
            return _save_resume_response(app, current_document, current_sidecar.get("validation", {}),
                                         current_sidecar["working_docx_reference"],
                                         current_sidecar["working_pdf_reference"],
                                         current_sidecar.get("saved_at", aa.now()), idempotent=True)

        resume_dir = ROOT / "output" / "resumes"
        reports_dir = ROOT / "output" / "reports"
        resume_dir.mkdir(parents=True, exist_ok=True)
        stamp = uuid.uuid4().hex[:10]
        stem = aa.artifact_stem(app, "resume", "Working", f"{candidate['revision_id']}_{stamp}")
        final_docx = resume_dir / f"{stem}.docx"
        final_pdf = resume_dir / f"{stem}.pdf"
        report_path = reports_dir / f"{aid}_resume_document_validation_{candidate['revision_id']}_{stamp}.json"
        sidecar_path = ROOT / "output" / "resume_edits" / aid / f"{candidate['revision_id']}_{stamp}.json"
        staged_paths: list[Path] = []

        try:
            with tempfile.TemporaryDirectory(prefix=f".{aid}_resume_edit_", dir=str(resume_dir)) as temporary:
                stage = Path(temporary)
                staged_docx, staged_pdf = stage / "candidate.docx", stage / "candidate.pdf"
                rendered = rdr.render_resume_document(candidate, validation_app, context["plan"], context["profile"],
                                                      staged_docx, staged_pdf, stage / "convert")
                validation_document = _frozen_generator_validation_view(candidate, baseline_document)
                validation_docx, validation_pdf = stage / "generator-validation-view.docx", stage / "generator-validation-view.pdf"
                rdr.render_resume_document(validation_document, validation_app, context["plan"], context["profile"],
                                           validation_docx, validation_pdf, stage / "generator-validation-convert")
                generated_report_path = stage / "frozen-validation.json"
                frozen_report = rg.validate(validation_docx, context["plan"], context["profile"], generated_report_path)
                frozen_report["validation_view"] = "temporary canonical labels for frozen-generator compatibility; Working DOCX/PDF are rendered from the unchanged user candidate"
                frozen_report["working_artifact_page_count"] = rendered.get("page_count")
                pdf_text = rendered.get("extracted_text", "")
                expected_paragraphs = rdm.resume_document_paragraphs(candidate)
                missing_paragraphs = [text for text in expected_paragraphs
                                      if _normalized_document_text(text) not in _normalized_document_text(pdf_text)]
                ats = frozen_report.get("ats_validation", {})
                approved_projects, _, _ = rdm._approved_context(app, context["plan"], context["profile"])
                approved_project_names = {project["record_id"]: project["name"] for project in approved_projects}
                rendered_project_names = [approved_project_names.get(block.get("project_id"), "")
                                          for section in candidate.get("content", {}).get("sections", [])
                                          for block in section.get("blocks", []) if block.get("type") == "project_entry"]
                truth = frozen_report.get("truth_provenance_validation", {})
                truth_pass = all(value is True for key, value in truth.items() if key != "unsupported_claims_detected") and not truth.get("unsupported_claims_detected")
                structure = rendered.get("structure_validation", {})
                required_checks = {
                    "one_page": rendered.get("page_count") == 1 and frozen_report.get("page_count") == 1,
                    "existing_generator_validation": frozen_report.get("final_status") == "PASS",
                    "required_headings": ats.get("standard_headings_present") is True,
                    "single_column": ats.get("single_column") is True and structure.get("single_column") is True,
                    "no_tables_textboxes_drawings_or_images": ats.get("problematic_tables_or_textboxes") is False and all(structure.get(key) is True for key in ("no_tables", "no_text_boxes", "no_drawings_images_or_floating_objects")),
                    "contact": frozen_report.get("contact_validation", {}).get("passed") is True,
                    "hyperlinks": frozen_report.get("hyperlink_validation", {}).get("passed") is True,
                    "project_links": frozen_report.get("project_link_validation", {}).get("passed") is True,
                    "project_order": frozen_report.get("project_order_validation", {}).get("passed") is True,
                    "selected_projects": rendered_project_names == frozen_report.get("selected_projects"),
                    "project_stacks": frozen_report.get("project_stack_validation", {}).get("passed") is True,
                    "skills": frozen_report.get("skills_validation", {}).get("passed") is True,
                    "summary": frozen_report.get("summary_validation", {}).get("passed") is True and frozen_report.get("summary_format_validation", {}).get("passed") is True,
                    "languages": frozen_report.get("languages_validation", {}).get("passed") is True,
                    "education": frozen_report.get("education_validation", {}).get("passed") is True,
                    "certifications": frozen_report.get("certification_validation", {}).get("passed") is True,
                    "placeholders": frozen_report.get("placeholder_validation", {}).get("passed") is True,
                    "extractable_text": frozen_report.get("ats_validation", {}).get("text_extractable") is True,
                    "truth_provenance": truth_pass,
                    "model_text_matches_pdf": not missing_paragraphs,
                }
                if not all(required_checks.values()):
                    print("RESUME_SAVE_VALIDATION_DEBUG", rdr.__file__, len(pdf_text), repr(pdf_text[:240]),
                          [repr(value) for value in expected_paragraphs[:3]], missing_paragraphs[:3], flush=True)
                    return {"decision": "validation_failed", "message": "The staged Working Resume did not pass all checks. The previous Working Resume remains active.",
                            "checks": required_checks, "missing_paragraphs": missing_paragraphs[:8],
                            "validation": frozen_report, "previous_working_preserved": True}

                saved_at = aa.now()
                docx_ref, pdf_ref = str(final_docx.relative_to(ROOT)), str(final_pdf.relative_to(ROOT))
                docx_hash, pdf_hash = file_sha256(staged_docx), file_sha256(staged_pdf)
                report = copy.deepcopy(frozen_report)
                report.update({"application_id": aid, "company_name": app.get("company_name"),
                               "revision_id": candidate["revision_id"],
                               "source_plan_sha256": candidate.get("source_plan_sha256"),
                               "source_generation_reference": candidate.get("source_generation_reference"),
                               "working_docx_reference": docx_ref, "working_pdf_reference": pdf_ref,
                               "working_docx_sha256": docx_hash, "working_pdf_sha256": pdf_hash,
                               "claim_validation": policy_report, "model_validation": model_report,
                               "checks": required_checks, "saved_at": saved_at})
                sidecar = {"sidecar_schema_version": 1, "application_id": aid,
                           "revision_id": candidate["revision_id"], "document": candidate,
                           "saved_at": saved_at, "working_docx_reference": docx_ref,
                           "working_pdf_reference": pdf_ref, "working_docx_sha256": docx_hash,
                           "working_pdf_sha256": pdf_hash, "source_plan_sha256": candidate.get("source_plan_sha256"),
                           "source_generation_reference": candidate.get("source_generation_reference"),
                           "validation_reference": str(report_path.relative_to(ROOT)),
                           "change_ledger": _build_resume_change_ledger(current_document, candidate),
                           "validation": report}

                # All checks are complete. Stage same-volume copies and activate the
                # application references last, with rollback for ordinary I/O errors.
                if final_docx.exists() or final_pdf.exists() or report_path.exists() or sidecar_path.exists():
                    raise FileExistsError("Revision destination already exists; refusing to overwrite an artifact.")
                shutil.copy2(staged_docx, final_docx); staged_paths.append(final_docx)
                shutil.copy2(staged_pdf, final_pdf); staged_paths.append(final_pdf)
                _atomic_write_json(report_path, report, exclusive=True); staged_paths.append(report_path)
                _atomic_write_json(sidecar_path, sidecar, exclusive=True); staged_paths.append(sidecar_path)

                latest_store = aa.load_store()
                if store_path.read_bytes() != original_store_bytes:
                    raise RuntimeError("Application data changed during validation. Reload the current application and retry.")
                target = next((item for item in latest_store.get("applications", []) if item.get("application_id") == aid), None)
                if not target:
                    raise RuntimeError("The selected application was removed during validation.")
                target.update({
                    "working_resume_reference": docx_ref,
                    "working_resume_pdf_reference": pdf_ref,
                    "working_resume_docx_path": docx_ref,
                    "working_resume_pdf_path": pdf_ref,
                    "working_resume_docx_sha256": docx_hash,
                    "working_resume_pdf_sha256": pdf_hash,
                    "working_resume_generation_id": validation_app["working_resume_generation_id"],
                    "working_resume_generated_at": saved_at,
                    "working_resume_revision_id": candidate["revision_id"],
                    "working_resume_revision_reference": str(sidecar_path.relative_to(ROOT)),
                    "working_resume_source_plan_sha256": candidate.get("source_plan_sha256"),
                    "working_resume_validated": True,
                    "working_resume_validated_at": saved_at,
                    "working_resume_validation_reference": str(report_path.relative_to(ROOT)),
                    "resume_validation_reference": str(report_path.relative_to(ROOT)),
                    "resume_working_artifact_stale": False,
                    "last_updated": saved_at,
                })
                _write_application_store_atomically(original_store_bytes, latest_store)
                staged_paths.clear()
                return _save_resume_response(target, candidate, report, docx_ref, pdf_ref, saved_at)
        except rdm.ResumeDocumentValidationError as exc:
            decision, message = "invalid", str(exc)
        except (OSError, RuntimeError, ValueError, KeyError, TypeError) as exc:
            decision, message = "error", str(exc)
        finally:
            for path in reversed(staged_paths):
                try:
                    path.unlink(missing_ok=True)
                except OSError:
                    pass
        return {"decision": decision, "message": f"Working Resume save was not activated: {message}",
                "previous_working_preserved": True}


def convert_pdf(docx):
    with tempfile.TemporaryDirectory() as temp_dir:
        _, _, pdf = rg.render_page_count(docx, Path(temp_dir))
        destination = docx.with_suffix(".pdf")
        shutil.copy2(pdf, destination)
    return destination


def persist_working_cover_letter_formats(aid, result):
    app = get_application(aid)
    if not app:
        return {"decision": "not_found"}
    markdown_ref = app.get("cover_letter_working_reference") or result.get("cover_letter_working_reference")
    if not markdown_ref:
        return {"decision": "error", "message": "The Working Cover Letter source was not saved."}
    markdown_path = resolve_ref(markdown_ref)
    if not markdown_path.is_file():
        return {"decision": "error", "message": "The Working Cover Letter source could not be found."}
    docx_path = markdown_path.with_suffix(".docx")
    from docx import Document
    document = Document()
    for line in markdown_path.read_text(encoding="utf-8").splitlines():
        document.add_paragraph(line)
    document.save(docx_path)
    pdf_path = convert_pdf(docx_path)
    if not pdf_path.is_file() or pdf_path.read_bytes()[:4] != b"%PDF":
        return {"decision": "error", "message": "The Working Cover Letter PDF could not be generated."}

    store = aa.load_store()
    target = next((item for item in store.get("applications", []) if item.get("application_id") == aid), None)
    if not target:
        return {"decision": "not_found"}
    docx_ref = str(docx_path.relative_to(ROOT))
    pdf_ref = str(pdf_path.relative_to(ROOT))
    target.update({
        "cover_letter_working_reference": str(markdown_ref),
        "cover_letter_source_reference": str(markdown_ref),
        "cover_letter_working_docx_reference": docx_ref,
        "cover_letter_working_pdf_reference": pdf_ref,
        "last_updated": aa.now(),
    })
    aa.save_store(store)
    return {
        **result,
        "cover_letter_working_reference": str(markdown_ref),
        "cover_letter_working_docx_reference": docx_ref,
        "cover_letter_working_pdf_reference": pdf_ref,
    }


def generate_working_cover_letter(aid):
    result = aa.generate_cover_letter(aid)
    if result.get("decision") != "created":
        return result
    return persist_working_cover_letter_formats(aid, result)


def generate_working_resume(aid):
    if rg is None:
        return {"decision": "backend_unavailable", "message": "Resume document tooling is unavailable in this runtime; use the project desktop service."}
    app = get_application(aid)
    if not app:
        return {"decision": "not_found"}
    if not app.get("resume_generation_allowed"):
        return {"decision": "approval_required", "message": "Approve the Resume Plan in JD Intelligence before generating a resume."}

    plan_path = resolve_ref(app.get("phase8_plan_reference"))
    missing_plan = not plan_path or not plan_path.exists()
    if missing_plan:
        jd_text = str(app.get("job_description_text") or "").strip()
        if not jd_text:
            return {"decision": "error", "message": "The approved Resume Plan could not be found and no durable JD text is available to reconstruct it."}

        plan = planner.plan_resume(jd_text, planner.load_profile())
    else:
        plan = json.loads(plan_path.read_text(encoding="utf-8"))

    if missing_plan or plan.get("approval_checkpoint", {}).get("resume_generation_allowed") is not True:
        plan.setdefault("approval_checkpoint", {})["resume_generation_allowed"] = True
        plan_payload = json.dumps(plan, indent=2, ensure_ascii=False) + "\n"
        plan_digest = hashlib.sha256(plan_payload.encode("utf-8")).hexdigest()[:16]
        plan_path = aa.REPORTS / f"{aid}_phase8_plan_{plan_digest}.json"
        plan_path.parent.mkdir(parents=True, exist_ok=True)
        if not plan_path.exists():
            plan_path.write_text(plan_payload, encoding="utf-8")

        store = aa.load_store()
        target = next(item for item in store["applications"] if item.get("application_id") == aid)
        target["phase8_plan_reference"] = aa.storage_reference(plan_path, root=aa.ROOT)
        aa.save_store(store)

    profile = rg.profile()
    output = aa.RESUMES / f"{aa.slug(app.get('company_name') or 'company')}_{aa.slug(app.get('job_title') or 'role')}_{aid}_Working.docx"
    report = aa.REPORTS / f"{aid}_resume_validation.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    report.parent.mkdir(parents=True, exist_ok=True)
    rg.generate(plan, profile, output)
    validation = rg.validate(output, plan, profile, report)
    if validation.get("page_count") != 1 or validation.get("final_status") != "PASS":
        return {"decision": "validation_failed", "message": "Resume validation did not pass; the working artifact was not activated.", "validation": validation, "validation_reference": str(report.relative_to(ROOT))}
    working_pdf = convert_pdf(output)
    generation_id = "gen_" + uuid.uuid4().hex
    generated_at = aa.now()
    docx_ref = str(output.relative_to(ROOT))
    pdf_ref = str(working_pdf.relative_to(ROOT))
    store = aa.load_store()
    target = next(item for item in store["applications"] if item.get("application_id") == aid)
    target.update({"working_resume_reference": docx_ref, "working_resume_pdf_reference": pdf_ref, "working_resume_docx_path": docx_ref, "working_resume_pdf_path": pdf_ref, "working_resume_docx_sha256": file_sha256(output), "working_resume_pdf_sha256": file_sha256(working_pdf), "working_resume_generation_id": generation_id, "working_resume_generated_at": generated_at, "resume_validation_reference": str(report.relative_to(ROOT)), "resume_working_artifact_stale": False, "resume_final_stale": bool(target.get("resume_generation_id")), "current_status": "resume_ready", "last_updated": generated_at})
    aa.save_store(store)
    return {"decision": "created", "application": target, "working": {"docx_reference": docx_ref, "pdf_reference": pdf_ref, "generation_id": generation_id, "generated_at": generated_at, "validation_reference": str(report.relative_to(ROOT)), "validation": validation}}


def validate_working_revision_for_finalization(aid, app=None):
    """Read-only guard used by finalization; it never copies or changes artifacts."""
    app = app or get_application(aid)
    if not app:
        return {"decision": "not_found"}
    if not app.get("resume_generation_allowed"):
        return {"decision": "approval_required", "message": "Approve the current Resume Plan before finalizing."}
    if app.get("resume_working_artifact_stale"):
        return {"decision": "stale_working", "message": "The Working Resume is stale because its plan or project selection changed. Approve and regenerate before finalizing."}
    working_ref = app.get("working_resume_docx_path") or app.get("working_resume_reference")
    pdf_ref = app.get("working_resume_pdf_path") or app.get("working_resume_pdf_reference")
    if not working_ref or not pdf_ref:
        return {"decision": "working_resume_required", "message": "A validated Working DOCX and PDF are required before finalizing."}
    source, pdf_path = _safe_project_path(working_ref), _safe_project_path(pdf_ref)
    if not source or not pdf_path or not source.is_file() or not pdf_path.is_file():
        return {"decision": "artifact_missing", "message": "The active Working Resume DOCX/PDF could not be safely resolved."}
    expected_docx_hash, expected_pdf_hash = app.get("working_resume_docx_sha256"), app.get("working_resume_pdf_sha256")
    if not expected_docx_hash or not expected_pdf_hash or file_sha256(source) != expected_docx_hash or file_sha256(pdf_path) != expected_pdf_hash:
        return {"decision": "artifact_modified", "message": "The active Working Resume no longer matches its validated SHA-256 metadata."}

    context = _resume_document_context(app)
    if context.get("decision") != "ready":
        decision = context.get("decision", "validation_failed")
        return {"decision": decision, "message": context.get("error", "The current Working Resume revision could not be validated.")}

    sidecar = context.get("sidecar")
    if sidecar:
        import resume_document_model as rdm
        import resume_document_validation as rdv
        document = context["document"]
        if (sidecar.get("application_id") != aid or sidecar.get("revision_id") != app.get("working_resume_revision_id")
                or document.get("revision_id") != app.get("working_resume_revision_id")
                or document.get("revision_id") != rdm._revision_id(document)):
            return {"decision": "revision_mismatch", "message": "The active Working Resume does not match the application's current revision ID."}
        if document.get("source_plan_sha256") != rdm._plan_digest(context["plan"]):
            return {"decision": "stale_working", "message": "The approved plan changed after this edited Working Resume was validated."}
        if document.get("source_generation_reference") != app.get("working_resume_generation_id"):
            return {"decision": "revision_mismatch", "message": "The edited revision is not associated with the current Working Resume generation."}
        if (sidecar.get("working_docx_reference") != working_ref or sidecar.get("working_pdf_reference") != pdf_ref
                or sidecar.get("working_docx_sha256") != expected_docx_hash
                or sidecar.get("working_pdf_sha256") != expected_pdf_hash):
            return {"decision": "revision_mismatch", "message": "The active artifact paths/hashes do not match the saved revision metadata."}
        validation_ref = app.get("working_resume_validation_reference")
        if not app.get("working_resume_validated") or validation_ref != sidecar.get("validation_reference"):
            return {"decision": "validation_report_missing", "message": "The active edited revision has no current passing validation reference."}
        if sidecar.get("validation", {}).get("final_status") != "PASS":
            return {"decision": "validation_failed", "message": "The saved revision's validation report is not PASS."}
        validation_path = _safe_project_path(validation_ref)
        if not validation_path or not validation_path.is_file():
            return {"decision": "validation_report_missing", "message": "The saved revision's validation report is missing."}
        try:
            persisted_report = json.loads(validation_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return {"decision": "validation_report_missing", "message": "The saved revision's validation report could not be read."}
        if (persisted_report.get("final_status") != "PASS" or persisted_report.get("application_id") != aid
                or persisted_report.get("revision_id") != document.get("revision_id")
                or persisted_report.get("working_docx_sha256") != expected_docx_hash
                or persisted_report.get("working_pdf_sha256") != expected_pdf_hash):
            return {"decision": "validation_report_missing", "message": "The stored validation report does not match the active application/revision/hashes."}
        policy = rdv.validate_edit_policy(document, context["baseline_document"], context["profile"])
        if not policy.get("valid"):
            return {"decision": "validation_failed", "message": "The edited revision failed current source-evidence checks."}
        return {"decision": "ready", "application": app, "source": source, "pdf": pdf_path,
                "plan": context["plan"], "profile": context["profile"],
                "revision_id": document["revision_id"], "sidecar": sidecar}

    # A generator-produced Working artifact still requires a current generation,
    # recorded PASS report, and byte-for-byte agreement with its app metadata.
    generation_id = app.get("working_resume_generation_id")
    if not isinstance(generation_id, str) or not generation_id.startswith("gen_"):
        return {"decision": "revision_mismatch", "message": "The Working Resume has no current generation/revision identifier."}
    if "_Working_rev_" in Path(working_ref).name:
        return {"decision": "revision_mismatch", "message": "An edited Working artifact is missing its matching revision sidecar."}
    validation_ref = app.get("resume_validation_reference")
    validation_path = _safe_project_path(validation_ref) if validation_ref else None
    if not validation_path or not validation_path.is_file():
        return {"decision": "validation_report_missing", "message": "The current Working Resume validation report is missing."}
    try:
        persisted_report = json.loads(validation_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"decision": "validation_report_missing", "message": "The current Working Resume validation report could not be read."}
    if persisted_report.get("final_status") != "PASS":
        return {"decision": "validation_failed", "message": "The current Working Resume has no passing validation report."}
    return {"decision": "ready", "application": app, "source": source, "pdf": pdf_path,
            "plan": context["plan"], "profile": context["profile"], "revision_id": generation_id, "sidecar": None}


def finalize_resume(aid):
    if rg is None:
        return {"decision": "backend_unavailable", "message": "Resume document tooling is unavailable in this runtime; use the project desktop service."}
    app = get_application(aid)
    if not app:
        return {"decision": "not_found"}
    guard = validate_working_revision_for_finalization(aid, app)
    if guard.get("decision") != "ready":
        return guard
    source, plan, profile = guard["source"], guard["plan"], guard["profile"]
    validation_report = aa.REPORTS / f"{aid}_finalization_validation.json"
    validation = rg.validate(source, plan, profile, validation_report)
    if validation.get("final_status") != "PASS":
        return {"decision": "validation_failed", "message": "Finalization blocked because the Working Resume did not pass validation.", "validation": validation, "validation_reference": str(validation_report.relative_to(ROOT))}
    final = source.with_name(source.stem.replace("_Working", "_Final") + source.suffix)
    if final.exists() or final.with_suffix(".pdf").exists():
        base = final.with_suffix("")
        version = 2
        while True:
            candidate = base.with_name(f"{base.name}_v{version}").with_suffix(final.suffix)
            if not candidate.exists() and not candidate.with_suffix(".pdf").exists():
                final = candidate
                break
            version += 1
    shutil.copy2(source, final)
    final_pdf = convert_pdf(final)
    final_docx_ref = str(final.relative_to(ROOT))
    final_pdf_ref = str(final_pdf.relative_to(ROOT))
    finalized_at = aa.now()
    store = aa.load_store()
    target = next(item for item in store["applications"] if item.get("application_id") == aid)
    target.update({"resume_reference": final_docx_ref, "resume_pdf_reference": final_pdf_ref, "resume_docx_path": final_docx_ref, "resume_pdf_path": final_pdf_ref, "resume_docx_sha256": file_sha256(final), "resume_pdf_sha256": file_sha256(final_pdf), "resume_generation_id": "gen_" + uuid.uuid4().hex, "resume_finalized_at": finalized_at, "resume_status": "final", "resume_final_stale": False, "last_updated": finalized_at})
    aa.save_store(store)
    return {"decision": "finalized", "application": target, "final": {"docx_reference": final_docx_ref, "pdf_reference": final_pdf_ref, "finalized_at": finalized_at, "validation_reference": str(validation_report.relative_to(ROOT))}}


PROJECT_LIMIT = 3


def eligible_completed_projects():
    projects = planner.load_profile().get("projects", {}).get("projects", [])
    seen = set()
    result = []
    for project in projects:
        record_id = project.get("record_id")
        if record_id and record_id not in seen and project.get("project_status") == "completed" and project.get("status") == "verified":
            result.append(project)
            seen.add(record_id)
    return result


def resume_editor_payload(aid):
    app = get_application(aid)
    if not app:
        return {"decision": "not_found"}
    plan_path = resolve_ref(app.get("phase8_plan_reference"))
    if not plan_path.exists():
        return {"decision": "error", "message": "The approved Resume Plan could not be found."}
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    resume_plan = plan.get("resume_plan", {})
    selected_ids = list(app.get("project_selection_record_ids") or [item.get("record_id") for item in resume_plan.get("projects_to_include", [])])
    automatic = resume_plan.get("automatic_projects_to_include") or resume_plan.get("projects_to_include", [])
    return {"decision": "ready", "application_id": aid, "mode": app.get("project_selection_mode") or resume_plan.get("project_selection_source") or "automatic", "selected_record_ids": selected_ids, "automatic_record_ids": [item.get("record_id") for item in automatic], "max_projects": PROJECT_LIMIT, "projects": [{"record_id": item.get("record_id"), "name": item.get("name"), "status": item.get("project_status"), "technologies": item.get("technologies") or item.get("frameworks_libraries_tools") or []} for item in eligible_completed_projects()], "limitations": ["Only canonical completed projects are editable.", "Project selection changes invalidate the current Working Resume and require plan approval before regeneration.", "Final Resume artifacts remain unchanged and protected."]}


def save_resume_edit(aid, payload):
    app = get_application(aid)
    if not app:
        return {"decision": "not_found"}
    mode = payload.get("mode", "manual")
    record_ids = list(payload.get("record_ids") or [])
    if mode not in {"automatic", "manual"}:
        return {"decision": "invalid", "message": "Invalid project-selection mode."}
    if len(record_ids) != len(set(record_ids)):
        return {"decision": "invalid", "message": "Duplicate projects are not allowed."}
    if len(record_ids) > PROJECT_LIMIT:
        return {"decision": "invalid", "message": f"At most {PROJECT_LIMIT} completed projects can be selected."}
    plan_path = resolve_ref(app.get("phase8_plan_reference"))
    if not plan_path.exists():
        return {"decision": "error", "message": "The Resume Plan could not be found."}
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    resume_plan = plan.setdefault("resume_plan", {})
    eligible = {item.get("record_id"): item for item in eligible_completed_projects()}
    if mode == "automatic":
        automatic = resume_plan.get("automatic_projects_to_include") or resume_plan.get("projects_to_include", [])
        record_ids = [item.get("record_id") for item in automatic if item.get("record_id") in eligible]
        selected = [eligible[item] for item in record_ids]
    else:
        if not record_ids:
            return {"decision": "invalid", "message": "Select at least one canonical completed project."}
        invalid = [item for item in record_ids if item not in eligible]
        if invalid:
            return {"decision": "invalid", "message": "Only canonical completed projects may be selected: " + ", ".join(invalid)}
        selected = [eligible[item] for item in record_ids]
    resume_plan["projects_to_include"] = selected
    resume_plan["project_selection_source"] = mode
    resume_plan["project_selection_record_ids"] = list(record_ids)
    plan["project_selection_source"] = mode
    plan.setdefault("approval_checkpoint", {})["resume_generation_allowed"] = False
    plan_path.write_text(json.dumps(plan, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    store = aa.load_store()
    target = next(item for item in store["applications"] if item.get("application_id") == aid)
    updated = aa.now()
    target.update({"selected_projects": [item.get("name") for item in selected], "project_selection_mode": mode, "project_selection_source": mode, "project_selection_record_ids": list(record_ids), "resume_generation_allowed": False, "resume_working_artifact_stale": True, "current_status": "awaiting_resume_approval", "last_updated": updated})
    aa.save_store(store)
    return {"decision": "saved", "message": "Project selection saved. Approve the updated Resume Plan in JD Intelligence, then regenerate the Working Resume. Final artifacts were not changed.", "application": target, "selected_record_ids": record_ids, "final_preserved": True}


class Handler(BaseHTTPRequestHandler):
    def _single_header(self, name, *, required=False):
        values = self.headers.get_all(name, [])
        if len(values) != 1:
            return None if not values and not required else False
        value = values[0]
        return value if value and "\r" not in value and "\n" not in value else False

    def _guard_request(self, *, require_json=False):
        host = self._single_header("Host", required=True)
        if host is False or host not in _ALLOWED_HOSTS:
            self.send_json({"error": "request denied"}, 403)
            return False

        origins = self.headers.get_all("Origin", [])
        if len(origins) > 1 or (origins and origins[0] not in _ALLOWED_ORIGINS):
            self.send_json({"error": "request denied"}, 403)
            return False

        request_path = urlparse(self.path).path
        if (request_path.startswith("/api/")
                and not _valid_api_token(self.command, request_path, self._single_header("Authorization"))):
            self.send_json({"error": "authentication required"}, 401)
            return False

        fetch_values = {
            "Sec-Fetch-Site": {"same-origin", "same-site", "none"},
            "Sec-Fetch-Mode": {"navigate", "same-origin", "no-cors", "cors", "websocket"},
            "Sec-Fetch-Dest": {
                "audio", "audioworklet", "document", "embed", "empty", "font", "frame",
                "fencedframe", "iframe", "image", "json", "manifest", "object", "paintworklet",
                "report", "script", "serviceworker", "sharedworker", "speculationrules", "style",
                "track", "video", "webidentity", "worker", "xslt",
            },
            "Sec-Fetch-User": {"?1"},
        }
        for name, allowed in fetch_values.items():
            values = self.headers.get_all(name, [])
            if len(values) > 1 or (values and values[0] not in allowed):
                self.send_json({"error": "request denied"}, 403)
                return False

        if require_json:
            content_type = self._single_header("Content-Type", required=True)
            if content_type is False or not re.fullmatch(
                r'application/json[ \t]*(?:;[ \t]*charset[ \t]*=[ \t]*(?:"utf-8"|utf-8)[ \t]*)?',
                content_type,
                re.IGNORECASE,
            ):
                self.send_json({"error": "application/json required"}, 415)
                return False
        return True

    def _send_cors_headers(self):
        origins = self.headers.get_all("Origin", [])
        if len(origins) == 1 and origins[0] in _ALLOWED_ORIGINS:
            self.send_header("Access-Control-Allow-Origin", origins[0])
            self.send_header("Vary", "Origin")

    def send_json(self, payload, status=200):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self._send_cors_headers()
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(body)

    def read_json(self):
        length = int(self.headers.get("Content-Length", "0"))
        if not length:
            return {}
        try:
            return json.loads(self.rfile.read(length).decode("utf-8"))
        except json.JSONDecodeError:
            return {}

    def do_OPTIONS(self):
        if not self._guard_request():
            return
        requested_methods = self.headers.get_all("Access-Control-Request-Method", [])
        if len(requested_methods) > 1:
            return self.send_json({"error": "request denied"}, 400)
        if requested_methods and requested_methods[0] not in {"GET", "POST", "OPTIONS"}:
            return self.send_json({"error": "method not allowed"}, 405)
        requested_headers = self.headers.get_all("Access-Control-Request-Headers", [])
        if len(requested_headers) > 1:
            return self.send_json({"error": "request denied"}, 400)
        if requested_headers and not requested_methods:
            return self.send_json({"error": "request denied"}, 400)
        if requested_headers:
            names = [item.strip().lower() for item in requested_headers[0].split(",")]
            if names != ["content-type"]:
                return self.send_json({"error": "request denied"}, 400)
        self.send_response(204)
        self.send_header("Allow", "GET, POST, OPTIONS")
        self._send_cors_headers()
        origin = self.headers.get("Origin")
        if origin in _ALLOWED_ORIGINS:
            self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
            self.send_header("Access-Control-Allow-Headers", "Content-Type")
            self.send_header("Access-Control-Max-Age", "600")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()

    def do_GET(self):
        if not self._guard_request():
            return
        parsed = urlparse(self.path)
        if parsed.scheme or parsed.netloc or parsed.fragment:
            return self.send_json({"error": "not found"}, 404)
        if _BAD_PERCENT_ESCAPE.search(parsed.query):
            return self.send_json({"error": "not found"}, 404)
        path = _decode_url_path(parsed.path)
        if path is None:
            return self.send_json({"error": "not found"}, 404)
        try:
            query = parse_qs(parsed.query, encoding="utf-8", errors="strict")
        except (UnicodeDecodeError, ValueError):
            return self.send_json({"error": "not found"}, 404)
        try:
            if path == "/api/health":
                return self.send_json({"ok": True, "service": "career-os-api"})
            if path == "/api/bootstrap":
                return self.send_json({"profile": profile_summary(), **application_payload()})
            if path == "/api/profile":
                return self.send_json(profile_summary())
            if path == "/api/applications":
                return self.send_json(application_payload())
            if path == "/api/search":
                return self.send_json({"applications": aa.search((query.get("q") or [None])[0], (query.get("status") or [None])[0])})
            parts = path.strip("/").split("/")
            if len(parts) == 4 and parts[:2] == ["api", "applications"] and parts[3] == "resume-editor":
                payload = resume_editor_payload(parts[2])
                return self.send_json(payload, 404 if payload.get("decision") == "not_found" else 200)
            if len(parts) == 4 and parts[:2] == ["api", "applications"] and parts[3] == "resume-document":
                payload = resume_document_payload(parts[2])
                status = 404 if payload.get("decision") == "not_found" else 409 if payload.get("decision") in {"stale_revision", "artifact_modified", "artifact_missing"} else 400 if payload.get("decision") == "invalid" else 200
                return self.send_json(payload, status)
            if path == "/api/artifact":
                references = query.get("ref") or []
                reference = references[0] if len(references) == 1 else ""
                candidate = _artifact_path(reference)
                if candidate is None:
                    return self.send_json({"error": "artifact not found"}, 404)
                try:
                    body = candidate.read_bytes()
                except OSError:
                    return self.send_json({"error": "artifact not found"}, 404)
                self.send_response(200)
                self._send_cors_headers()
                self.send_header("Content-Type", mimetypes.guess_type(str(candidate))[0] or "application/octet-stream")
                self.send_header("Content-Length", str(len(body)))
                view_requested = (query.get("view") or [""])[0].lower() in {"1", "true", "inline"}
                disposition = "inline" if candidate.suffix.lower() == ".pdf" and view_requested else "attachment"
                self.send_header("Content-Disposition", f"{disposition}; filename=\"{candidate.name}\"")
                self.send_header("X-Content-Type-Options", "nosniff")
                self.end_headers()
                self.wfile.write(body)
                return
            return self.serve_frontend(parsed.path)
        except Exception as exc:
            return self.send_json({"error": f"{type(exc).__name__}: {exc}"}, 500)

    def do_POST(self):
        if not self._guard_request(require_json=True):
            return
        parsed = urlparse(self.path)
        if parsed.scheme or parsed.netloc or parsed.fragment:
            return self.send_json({"error": "not found"}, 404)
        if _BAD_PERCENT_ESCAPE.search(parsed.query):
            return self.send_json({"error": "not found"}, 404)
        path = _decode_url_path(parsed.path)
        if path is None:
            return self.send_json({"error": "not found"}, 404)
        payload = self.read_json()
        try:
            if path == "/api/analyze":
                jd = str(payload.get("job_description", "")).strip()
                if not jd:
                    return self.send_json({"error": "job_description is required"}, 400)
                profile = planner.load_profile()
                plan = planner.plan_resume(jd, profile)
                plan["candidate_skill_categories"] = [group.get("category") for group in profile.get("skills", {}).get("skill_groups", []) if group.get("category")]
                return self.send_json(plan)
            if path == "/api/applications":
                jd = str(payload.get("job_description", "")).strip()
                if not jd:
                    return self.send_json({"error": "job_description is required"}, 400)
                result = aa.create_application(payload.get("company"), payload.get("title"), payload.get("url"), jd, payload.get("source") or "user_provided", payload.get("location"), payload.get("employment_type"))
                status = 201 if result.get("decision") == "created" else 409 if result.get("decision", "").startswith("duplicate") else 400
                return self.send_json(result, status)
            parts = path.strip("/").split("/")
            if len(parts) == 5 and parts[:2] == ["api", "applications"] and parts[3:] == ["resume-document", "save"]:
                result = save_resume_document_payload(parts[2], payload)
                status = 404 if result.get("decision") == "not_found" else 409 if result.get("decision") in {"approval_required", "stale_revision", "validation_failed", "artifact_modified", "artifact_missing", "stale_working"} else 500 if result.get("decision") == "error" else 400 if result.get("decision") in {"invalid", "backend_unavailable"} else 200
                return self.send_json(result, status)
            if len(parts) == 5 and parts[:2] == ["api", "applications"] and parts[3:] == ["resume-document", "validate"]:
                result = validate_resume_document_payload(parts[2], payload)
                status = 404 if result.get("decision") == "not_found" else 409 if result.get("decision") == "stale_revision" else 400 if result.get("decision") == "invalid" else 200
                return self.send_json(result, status)
            if len(parts) == 4 and parts[:2] == ["api", "applications"]:
                aid, action = parts[2], parts[3]
            else:
                return self.send_json({"error": "not found"}, 404)
            if action == "approve-resume":
                result = aa.approve_resume(aid, payload.get("plan"))
                status = 200 if result.get("decision") == "approved" else 404 if result.get("decision") == "not_found" else 400 if result.get("decision", "").startswith("plan_") or result.get("decision") == "invalid_plan" else 200
                return self.send_json(result, status)
            if action == "confirm-skill-gap":
                result = confirm_skill_gap(aid, payload)
                status = 200 if result.get("decision") == "skill_added_candidate_provided" else 404 if result.get("decision") == "not_found" else 400 if result.get("decision") in {"invalid", "confirmation_required", "no_change_duplicate", "ask_clarification"} else 500 if result.get("decision") == "error" else 200
                return self.send_json(result, status)
            if action == "resume-edit":
                result = save_resume_edit(aid, payload)
                return self.send_json(result, 400 if result.get("decision") in {"invalid", "error"} else 404 if result.get("decision") == "not_found" else 200)
            if action == "generate-resume":
                result = generate_working_resume(aid)
                return self.send_json(result, 409 if result.get("decision") in {"approval_required", "validation_failed"} else 404 if result.get("decision") == "not_found" else 200)
            if action == "finalize-resume":
                result = finalize_resume(aid)
                return self.send_json(result, 409 if result.get("decision") in {"working_resume_required", "artifact_missing", "validation_failed", "stale_working", "approval_required", "artifact_modified", "stale_revision", "revision_mismatch", "validation_report_missing"} else 404 if result.get("decision") == "not_found" else 200)
            if action == "manual-edit":
                return self.send_json({"decision": "unsupported", "message": "Manual resume editing is not exposed here because unsupported claims must remain governed by the existing Resume Project Edit/profile-evidence workflow."}, 409)
            if action == "cover-letter":
                result = generate_working_cover_letter(aid)
                status = 404 if result.get("decision") == "not_found" else 500 if result.get("decision") == "error" else 200
                return self.send_json(result, status)
            if action == "cover-letter-edit":
                result = aa.edit_cover_letter(aid, payload.get("content"))
                if result.get("decision") == "saved":
                    result = persist_working_cover_letter_formats(aid, result)
                status = 400 if result.get("decision") == "invalid" else 404 if result.get("decision") == "not_found" else 500 if result.get("decision") == "error" else 200
                return self.send_json(result, status)
            if action == "status":
                try:
                    result = aa.update_status(aid, payload.get("status"), bool(payload.get("explicit_submission")))
                    return self.send_json(result, 404 if result.get("decision") == "not_found" else 200)
                except ValueError as exc:
                    return self.send_json({"error": str(exc)}, 400)
            if action == "note":
                result = aa.add_note(aid, str(payload.get("note", "")), payload.get("follow_up_date"))
                return self.send_json(result, 404 if result.get("decision") == "not_found" else 200)
            return self.send_json({"error": "not found"}, 404)
        except PermissionError as exc:
            return self.send_json({"error": str(exc), "decision": "approval_required"}, 409)
        except Exception as exc:
            return self.send_json({"error": f"{type(exc).__name__}: {exc}"}, 500)

    def serve_frontend(self, path):
        decoded = _decode_url_path(path)
        if decoded is None:
            return self.send_json({"error": "not found"}, 404)
        if not FRONTEND.exists():
            return self.send_json({"error": "Frontend build not found. Run npm run build in frontend/."}, 503)
        relative = decoded[1:] or "index.html"
        candidate = _contained_file(FRONTEND, relative)
        if candidate is None:
            candidate = _contained_file(FRONTEND, "index.html")
        if candidate is None:
            return self.send_json({"error": "not found"}, 404)
        try:
            body = candidate.read_bytes()
        except OSError:
            return self.send_json({"error": "not found"}, 404)
        self.send_response(200)
        self._send_cors_headers()
        self.send_header("Content-Type", mimetypes.guess_type(str(candidate))[0] or "application/octet-stream")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format, *args):
        return


if __name__ == "__main__":
    if _RUNTIME_ENV in {"production", "prod"} and len(_API_TOKEN) < 32:
        raise SystemExit("CAREER_OS_API_TOKEN must be configured with at least 32 characters in production")
    ThreadingHTTPServer(_api_server_address(), Handler).serve_forever()
