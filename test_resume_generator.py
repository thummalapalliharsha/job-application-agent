#!/usr/bin/env python3
"""Focused tests for JD-tailored deterministic resume inputs."""
from __future__ import annotations

import copy
import json
import re
import tempfile
import unittest
from pathlib import Path

from docx import Document

import application_assistant as aa
import jd_resume_planner as planner
import resume_generator as rg
from tests.jd_regression_suite import run_suite as jd_suite


class PdfTextConsistencyNormalizationTests(unittest.TestCase):
    def assert_pdf_matches_docx(self, docx_text, pdf_text, layout_text):
        self.assertEqual(
            jd_suite._normalize(docx_text),
            jd_suite._normalize_pdf_line_end_hyphenation(pdf_text, docx_text, layout_text),
        )
        self.assertIsNone(jd_suite._first_text_difference(docx_text, pdf_text, layout_text))

    def test_calinski_harabasz_line_end_hyphenation(self):
        self.assert_pdf_matches_docx(
            "Calinski-Harabasz metrics",
            "calinskiharabasz metrics",
            "Calinski-\nHarabasz metrics",
        )

    def test_scikit_learn_line_end_hyphenation(self):
        self.assert_pdf_matches_docx(
            "Scikit-learn pipeline",
            "scikitlearn pipeline",
            "Scikit-\nlearn pipeline",
        )

    def test_ordinary_hyphenated_text_is_preserved(self):
        self.assert_pdf_matches_docx(
            "A data-driven workflow",
            "A data-driven workflow",
            "A data-driven workflow",
        )
        self.assertNotEqual(
            jd_suite._normalize("A data-driven workflow"),
            jd_suite._normalize_pdf_line_end_hyphenation(
                "A data driven workflow", "A data-driven workflow", "A data-driven workflow"
            ),
        )
        self.assertIsNotNone(
            jd_suite._first_text_difference(
                "A data-driven workflow", "A data driven workflow", "A data-driven workflow"
            )
        )

    def test_genuinely_missing_text_still_differs(self):
        docx_text = "Calinski-Harabasz metrics were evaluated."
        pdf_text = "Calinski metrics were evaluated."
        layout_text = "Calinski-\nHarabasz metrics were evaluated."
        self.assertNotEqual(
            jd_suite._normalize(docx_text),
            jd_suite._normalize_pdf_line_end_hyphenation(pdf_text, docx_text, layout_text),
        )
        self.assertIsNotNone(jd_suite._first_text_difference(docx_text, pdf_text, layout_text))

    def test_genuinely_different_text_still_differs(self):
        docx_text = "A model-based workflow was used."
        pdf_text = "A model driven workflow was used."
        self.assertNotEqual(
            jd_suite._normalize(docx_text),
            jd_suite._normalize_pdf_line_end_hyphenation(pdf_text, docx_text, ""),
        )
        self.assertIsNotNone(jd_suite._first_text_difference(docx_text, pdf_text))


class ResumeGeneratorSelectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.profile = rg.profile()

    def test_relevant_experience_is_not_suppressed_by_role_title(self):
        experience = self.profile["experience"]["experiences"][0]
        plan = {
            "jd_analysis": {"job_title": "Generative AI Engineer"},
            "resume_plan": {"experience_to_include": [{"record_id": experience["record_id"]}]},
        }
        self.assertEqual(rg.effective_experience(plan, self.profile, []), [experience])

    def test_only_verified_skills_and_approved_certifications_are_rendered(self):
        profile = copy.deepcopy(self.profile)
        profile["skills"]["skill_groups"][0]["skills"].append({"name": "Unverified Skill", "status": "candidate_provided"})
        certifications = profile["certifications"]["certifications"]
        verified = copy.deepcopy(certifications[0])
        verified.update({"record_id": "credential_verified_test", "status": "verified", "name": "Python Data Science Certification"})
        profile["certifications"]["certifications"].append(verified)
        candidate_provided_id = certifications[0]["record_id"]
        plan = {
            "source_jd_text": "Junior Data Analyst. Required Skills: Python, SQL, data analysis, reporting.",
            "jd_analysis": {"target_role": "Junior Data Analyst", "requirements": []},
            "resume_plan": {
            "skills_to_include": [
                {"name": "Python", "status": "verified"},
                {"name": "Unverified Skill", "status": "candidate_provided"},
            ],
            "certifications_to_include": [
                {"record_id": "credential_verified_test"},
                {"record_id": candidate_provided_id},
            ],
        }}
        self.assertEqual(rg.effective_skills(plan, profile), ["Python"])
        rendered = rg.effective_certs(plan, profile)
        self.assertEqual([item["record_id"] for item in rendered], ["credential_verified_test", candidate_provided_id])
        self.assertEqual([item["status"] for item in rendered], ["verified", "candidate_provided"])
        self.assertEqual(rg.effective_certs({"resume_plan": {"certifications_to_include": []}}, profile), [])

    def test_role_label_and_summary_are_specific_and_deterministic(self):
        profile = copy.deepcopy(self.profile)
        jd = "Junior Business Data Analyst — Fresher\nRequired Skills:\n- Python\n- SQL\nResponsibilities:\n- Prepare business reports."
        first = planner.plan_resume(jd, profile)
        second = planner.plan_resume(jd, profile)
        self.assertEqual(first["jd_analysis"]["target_role"], "Junior Business Data Analyst")
        self.assertEqual([x["record_id"] for x in first["resume_plan"]["projects_to_include"]],
                         [x["record_id"] for x in second["resume_plan"]["projects_to_include"]])
        self.assertEqual(rg.summary_lines(first, profile), rg.summary_lines(second, profile))
        self.assertIn("Junior Business Data Analyst", " ".join(rg.summary_lines(first, profile)))

    def test_summary_uses_data_analyst_evidence_instead_of_generic_phrase(self):
        profile = copy.deepcopy(self.profile)
        jd = "Junior Data Analyst\nRequired Skills:\n- Python\n- SQL\n- Pandas\n- NumPy\n- Data Visualization\nResponsibilities:\n- Clean and analyze structured datasets.\n- Create business-facing dashboards."
        plan = planner.plan_resume(jd, profile)
        summary = " ".join(rg.summary_lines(plan, profile))
        self.assertIn("Junior Data Analyst", summary)
        self.assertIn("Python", summary)
        self.assertIn("SQL", summary)
        self.assertIn("data cleaning", summary.lower())
        self.assertIn("dashboard", summary.lower())
        self.assertNotIn("This work reflects an analytical focus on exploratory analysis and clear business reporting.", summary)

    def test_summary_uses_ml_model_evidence_instead_of_generic_phrase(self):
        profile = copy.deepcopy(self.profile)
        jd = "Junior Machine Learning Engineer\nRequired Skills:\n- Python\n- Pandas\n- NumPy\n- Scikit-learn\n- Model Evaluation\nResponsibilities:\n- Train and compare models.\n- Evaluate performance and document experiments."
        plan = planner.plan_resume(jd, profile)
        summary = " ".join(rg.summary_lines(plan, profile))
        self.assertIn("Junior Machine Learning Engineer", summary)
        self.assertIn("scikit-learn", summary.lower())
        self.assertIn("model training", summary.lower())
        self.assertIn("evaluation", summary.lower())
        self.assertNotIn("This work reflects an analytical focus on exploratory analysis and clear business reporting.", summary)

    def test_project_bullets_come_from_canonical_functionality(self):
        project = next(item for item in self.profile["projects"]["projects"]
                       if item["record_id"] == "project_weather_app")
        bullets = rg.project_bullets(project)
        canonical = " ".join(project.get("functionality", []) + project.get("technical_details", []))
        self.assertTrue(bullets)
        self.assertTrue(all(bullet.rstrip(".") in canonical for bullet in bullets))
        self.assertFalse(any("18%" in bullet or "22%" in bullet for bullet in bullets))

    def test_regression_suite_project_bullets_preserve_content_and_quality(self):
        expected_selections = {
            "01_data_analyst": ["project_bank_customer_clustering_dashboard", "project_sample_sales_data", "project_telecom_churn_logistic_regression"],
            "02_business_data_analyst": ["project_bank_customer_clustering_dashboard", "project_sample_sales_data", "project_imdb_movie_analysis"],
            "03_machine_learning_engineer": ["project_smartfraud_classifier", "project_telecom_churn_logistic_regression", "project_fuel_regression_crispmlq"],
            "04_data_scientist": ["project_smartfraud_classifier", "project_telecom_churn_logistic_regression", "project_fuel_regression_crispmlq"],
            "05_rag_engineer": ["project_student_performance_rag", "project_text_to_sql_project", "project_ai_resume_jd_match_tool"],
            "06_ai_ml_engineer": ["project_smartfraud_classifier", "project_telecom_churn_logistic_regression", "project_fuel_regression_crispmlq"],
            "07_python_developer_data_ai": ["project_bank_customer_clustering_dashboard", "project_telecom_churn_logistic_regression", "project_fuel_regression_crispmlq"],
            "08_bi_data_visualization_analyst": ["project_sample_sales_data", "project_bank_customer_clustering_dashboard", "project_imdb_movie_analysis"],
            "09_nlp_llm_engineer": ["project_student_performance_rag", "project_text_to_sql_project", "project_ai_resume_jd_match_tool"],
            "10_data_engineer": ["project_student_performance_rag", "project_genai_pyspark_pipeline", "project_telecom_churn_logistic_regression"],
        }
        expected_reviewed_bullets = {
            "project_smartfraud_classifier": [
                "Loaded and explored the Kaggle fraud dataset for binary fraud classification.",
                "Compared Decision Tree, Random Forest, and XGBoost models for fraud detection.",
                "Applied SMOTE for class-imbalance handling and saved model artifacts.",
            ],
            "project_telecom_churn_logistic_regression": [
                "Inspected telecom customer data and performed exploratory data analysis.",
                "Built a binary churn-classification pipeline with scaling, one-hot encoding, and class-imbalance handling.",
                "Optimized logistic regression with GridSearchCV and provided churn risk scores through a Streamlit interface.",
            ],
            "project_fuel_regression_crispmlq": [
                "Prepared flight data with numerical and categorical preprocessing for fuel-consumption prediction.",
                "Used Linear Regression, Ridge, Lasso, and ElasticNet within the CRISP-ML(Q) workflow.",
                "Evaluated predictions with MAE, MSE, RMSE, MAPE, and R2 metrics through a Streamlit prediction interface.",
            ],
            "project_genai_pyspark_pipeline": [
                "Generated configurable customer, product, and order data with Faker.",
                "Exported synthetic datasets to Parquet and loaded them with Spark for analytics.",
                "Used PySpark aggregations to analyze customer revenue, category sales, monthly trends, and frequently purchased product combinations.",
            ],
        }
        supported_claims = {
            "project_smartfraud_classifier": ["Kaggle", "Decision Tree", "Random Forest", "XGBoost", "SMOTE"],
            "project_telecom_churn_logistic_regression": ["GridSearchCV", "Streamlit"],
            "project_fuel_regression_crispmlq": ["Linear Regression", "Ridge", "Lasso", "ElasticNet", "CRISP-ML(Q)", "MAE", "MSE", "RMSE", "MAPE", "R2", "Streamlit"],
            "project_genai_pyspark_pipeline": ["Faker", "Parquet", "Spark", "PySpark"],
        }
        cases_path = Path(__file__).resolve().parent / "tests" / "jd_regression_suite" / "expected_profiles.json"
        cases = json.loads(cases_path.read_text(encoding="utf-8"))["cases"]
        sentence_action = re.compile(
            r"\b(?:analyzed|applied|built|compared|created|displayed|evaluated|exported|generated|"
            r"inspected|integrated|loaded|optimized|performed|preprocessed|prepared|removed|returned|"
            r"saved|stored|used)\b",
            re.IGNORECASE,
        )
        with tempfile.TemporaryDirectory(prefix="project-bullet-quality-") as temporary:
            for case in cases:
                plan = planner.plan_resume(case["jd_text"], copy.deepcopy(self.profile))
                selected = rg.effective_selection(plan, self.profile)
                selected_ids = [item["record_id"] for item in selected]
                self.assertEqual(selected_ids, expected_selections[case["id"]], case["role"])
                plan["approval_checkpoint"]["resume_generation_allowed"] = True
                output = Path(temporary) / f"{case['id']}.docx"
                rendered_projects, _, _, _ = rg.generate(plan, self.profile, output)
                self.assertEqual([item["record_id"] for item in rendered_projects], expected_selections[case["id"]])

                paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
                project_start = paragraphs.index("PROJECTS")
                project_end = next(
                    index for index in range(project_start + 1, len(paragraphs))
                    if paragraphs[index] in {"RELEVANT EXPERIENCE", "EDUCATION"}
                )
                project_paragraphs = paragraphs[project_start + 1:project_end]
                expected_stacks = []
                expected_bullets = []
                for project in selected:
                    technologies = project.get("technologies") or project.get("frameworks_libraries_tools") or []
                    stack = list(dict.fromkeys(technologies))[:8]
                    expected_stacks.append("Tech Stack: " + ", ".join(stack))
                    bullets = rg.project_bullets(project, plan)[:3]
                    expected_bullets.extend(bullets)
                    self.assertTrue(
                        all(len(bullet.split()) >= 7 and bullet[0].isupper() and bullet.endswith((".", "!", "?"))
                            and sentence_action.search(bullet) for bullet in bullets),
                        f"Incomplete project bullet for {project['record_id']}: {bullets}",
                    )
                    if project["record_id"] in expected_reviewed_bullets:
                        self.assertEqual(bullets, expected_reviewed_bullets[project["record_id"]])
                        source = json.dumps(project, ensure_ascii=False).casefold()
                        for claim in supported_claims[project["record_id"]]:
                            self.assertIn(claim.casefold(), source)
                        source_numbers = {
                            re.sub(r"\D", "", number)
                            for number in re.findall(r"(?<!\w)\d[\d,]*(?:\.\d+)?%?(?!\w)", source)
                        }
                        for bullet in bullets:
                            for number in re.findall(r"(?<!\w)\d[\d,]*(?:\.\d+)?%?(?!\w)", bullet):
                                self.assertIn(re.sub(r"\D", "", number), source_numbers)

                actual_stacks = [line for line in project_paragraphs if line.startswith("Tech Stack: ")]
                actual_bullets = [line[2:] for line in project_paragraphs if line.startswith("• ")]
                self.assertEqual(actual_stacks, expected_stacks, case["role"])
                self.assertEqual(actual_bullets, expected_bullets, case["role"])

                if case["id"] == "05_rag_engineer":
                    self.assertEqual(rg.summary_lines(plan, self.profile), [
                        "Computer Science fresher targeting a Junior Generative AI / RAG Engineer role, applying Python to RAG, embeddings, and ChromaDB vector retrieval in the Student Performance RAG Chatbot.",
                        "The RAG workflow preprocesses student data into searchable profiles, retrieves relevant context, and uses local Ollama generation through Streamlit.",
                        "A supporting Text-to-SQL project uses the Gemini REST API to generate SQLite queries from natural-language questions.",
                    ])

    def test_insightedge_business_analyst_content_and_document_validation(self):
        application = next(item for item in aa.load_store()["applications"]
                           if item["application_id"] == "app_d003a3f71d8c")
        plan = planner.plan_resume(application["job_description_text"], copy.deepcopy(self.profile))
        projects = plan["resume_plan"]["projects_to_include"]
        project_ids = [project["record_id"] for project in projects]
        self.assertEqual(project_ids, [
            "project_bank_customer_clustering_dashboard",
            "project_sample_sales_data",
            "project_imdb_movie_analysis",
        ])
        certification_ids = [item["record_id"] for item in plan["resume_plan"]["certifications_to_include"]]
        self.assertEqual(certification_ids, [
            "credential_altair_rapidminer",
            "credential_eduskills_ai_ml_virtual_internship",
            "credential_ediglobe_ai_internship",
            "credential_ramp_coding",
        ])

        skill_groups = rg.effective_skill_groups(plan, self.profile, rg.effective_selection(plan, self.profile))
        self.assertEqual(skill_groups, [
            ("Programming", ["Python"]),
            ("Data Analysis", ["Pandas", "NumPy"]),
            ("Databases", ["SQL", "SQLite", "MySQL"]),
            ("Data Visualization", ["Matplotlib", "Seaborn", "Plotly"]),
            ("Tools", ["Streamlit"]),
        ])
        skill_text = " ".join(f"{label}: {', '.join(names)}" for label, names in skill_groups)
        self.assertNotIn("Text-to-SQL", skill_text)
        self.assertNotIn("Power BI", skill_text)
        self.assertNotIn("Excel", skill_text)

        expected_bullets = {
            "project_bank_customer_clustering_dashboard": [
                "Cleaned an 8,950-record credit-card dataset, engineered features, and applied robust standardization.",
                "Applied K-Means, Agglomerative, and DBSCAN clustering; evaluated clusters with Silhouette, Davies-Bouldin, and Calinski-Harabasz metrics.",
                "Built a Streamlit dashboard with Plotly visualizations for customer persona classification, PCA outputs, and density-based outlier detection.",
            ],
            "project_sample_sales_data": [
                "Generated synthetic transaction data with Pandas and NumPy, including product, category, customer, quantity, price, revenue, date, and region fields.",
                "Analyzed the generated data with D-Tale for interactive exploratory data analysis.",
                "Created HTML-based data profiling reports with YData Profiling and Sweetviz.",
            ],
            "project_imdb_movie_analysis": [
                "Loaded and inspected the IMDb movie dataset with Python, Pandas, and NumPy.",
                "Performed exploratory data analysis in a Jupyter notebook on the movie dataset.",
            ],
        }
        for project in rg.effective_selection(plan, self.profile):
            bullets = rg.project_bullets(project, plan)
            self.assertEqual(bullets, expected_bullets[project["record_id"]])
            self.assertTrue(all(bullet[0].isupper() and bullet.endswith(".") for bullet in bullets))
        movie_bullets = " ".join(expected_bullets["project_imdb_movie_analysis"]).casefold()
        for unsupported_claim in ("5,000+", "22%", "10+ visualizations", "director-impact"):
            self.assertNotIn(unsupported_claim, movie_bullets)

        summary = " ".join(rg.summary_lines(plan, self.profile))
        self.assertEqual(len(re.findall(r"(?<=[.!?])\s+", summary)), 2)
        self.assertLess(len(summary), 400)
        self.assertIn("Computer Science fresher", summary)
        self.assertIn("Junior Business Data Analyst", summary)
        self.assertIn("Python, SQL, Pandas, and NumPy", summary)
        self.assertIn("data cleaning", summary)
        self.assertIn("business reporting", summary)
        self.assertNotIn("Text-to-SQL", summary)
        self.assertTrue(all(bullet not in summary for bullets in expected_bullets.values() for bullet in bullets))

        plan["approval_checkpoint"]["resume_generation_allowed"] = True
        with tempfile.TemporaryDirectory(prefix="insightedge-resume-content-") as temporary:
            output = Path(temporary) / "insightedge_content_quality.docx"
            report_path = Path(temporary) / "validation.json"
            rg.generate(plan, self.profile, output)
            report = rg.validate(output, plan, self.profile, report_path)
            document = Document(output)
            paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
            rendered_text = "\n".join(paragraph.text for paragraph in paragraphs)

        self.assertEqual(report["final_status"], "PASS", report.get("warnings_issues"))
        self.assertTrue(report["page_count_exactly_one"])
        self.assertTrue(report["ats_validation"]["text_extractable"])
        self.assertTrue(report["ats_validation"]["single_column"])
        self.assertTrue(report["docx_structure_validation"]["no_tables"])
        self.assertTrue(report["docx_structure_validation"]["no_text_boxes"])
        self.assertTrue(report["docx_structure_validation"]["no_graphics"])
        self.assertEqual(report["selected_projects"], [project["name"] for project in projects])
        self.assertEqual(report["selected_certifications"], [item["name"] for item in plan["resume_plan"]["certifications_to_include"]])
        self.assertNotIn("RELEVANT EXPERIENCE", rendered_text)
        self.assertNotIn("Generative AI: Text-to-SQL", rendered_text)
        self.assertIn("Data Analysis: Pandas, NumPy", rendered_text)
        self.assertIn("Data Visualization: Matplotlib, Seaborn, Plotly", rendered_text)

        paragraph_by_text = {paragraph.text: paragraph for paragraph in paragraphs}
        for heading in ("PROFESSIONAL SUMMARY", "SKILLS", "PROJECTS", "EDUCATION", "CERTIFICATIONS"):
            self.assertEqual(paragraph_by_text[heading].paragraph_format.space_before.pt, 5)
        project_titles = [paragraph_by_text[project["name"]] for project in projects]
        self.assertEqual(project_titles[0].paragraph_format.space_before.pt, 1.5)
        self.assertEqual(project_titles[0].paragraph_format.space_after.pt, 1.5)
        self.assertEqual([title.paragraph_format.space_before.pt for title in project_titles[1:]], [5.5, 5.5])
        for project in projects:
            project_bullets = [paragraph for paragraph in paragraphs if paragraph.text.startswith("• ") and paragraph.text[2:] in expected_bullets[project["record_id"]]]
            self.assertEqual(len(project_bullets), len(expected_bullets[project["record_id"]]))
            self.assertTrue(all(bullet.paragraph_format.space_after.pt == 1.5 for bullet in project_bullets))
            title_index = paragraphs.index(paragraph_by_text[project["name"]])
            following = paragraphs[title_index + 1:]
            tech_stack = next(paragraph for paragraph in following if paragraph.text.startswith("Tech Stack:"))
            project_link = next(paragraph for paragraph in following if paragraph.text.startswith("Project Link:"))
            self.assertEqual(tech_stack.paragraph_format.space_before.pt, 1.5)
            self.assertEqual(tech_stack.paragraph_format.space_after.pt, 1.5)
            self.assertEqual(project_link.paragraph_format.space_before.pt, 1.5)
            self.assertEqual(project_link.paragraph_format.space_after.pt, 1.5)
        education_lines = [paragraph for paragraph in paragraphs if "Class XII" in paragraph.text or "Class X —" in paragraph.text]
        self.assertEqual(len(education_lines), 2)
        self.assertTrue(all(paragraph.paragraph_format.space_before.pt == 2.5 for paragraph in education_lines))
        certification_start = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text == "CERTIFICATIONS")
        certification_bullets = [paragraph for paragraph in paragraphs[certification_start + 1:] if paragraph.text.startswith("• ")]
        self.assertEqual(len(certification_bullets), 4)
        self.assertTrue(all(paragraph.paragraph_format.space_after.pt == 2.5 for paragraph in certification_bullets))

    def test_plain_jd_skill_lists_are_classified_by_section(self):
        jd = """Junior Engineer
Responsibilities
Build reliable data tools.
Required Skills
Python
RAG
Preferred Skills
ChromaDB
Education
Bachelor's degree in Computer Science.
"""
        requirements = planner.analyze_jd(jd)["requirements"]
        classifications = {item["canonical"]: item["classification"] for item in requirements}
        self.assertEqual(classifications["python"], "required")
        self.assertEqual(classifications["rag"], "required")
        self.assertEqual(classifications["chromadb"], "preferred")

    def test_vectormind_rag_categories_jd_coverage_and_generation(self):
        application = next(item for item in aa.load_store()["applications"]
                           if item["application_id"] == "app_03fefedd7498")
        plan = planner.plan_resume(application["job_description_text"], copy.deepcopy(self.profile))
        projects = plan["resume_plan"]["projects_to_include"]
        self.assertEqual([item["record_id"] for item in projects], [
            "project_student_performance_rag",
            "project_text_to_sql_project",
            "project_ai_resume_jd_match_tool",
        ])
        certifications = plan["resume_plan"]["certifications_to_include"]
        self.assertEqual([item["record_id"] for item in certifications], [
            "credential_altair_rapidminer",
            "credential_eduskills_ai_ml_virtual_internship",
            "credential_ediglobe_ai_internship",
            "credential_ramp_coding",
        ])
        expected_project_bullets = {
            "project_student_performance_rag": [
                "Preprocessed student-performance data and created searchable student profiles for a local RAG chatbot.",
                "Generated local embeddings, stored student profiles in ChromaDB, and retrieved relevant records for responses from a local Ollama model.",
                "Displayed retrieved sources in Streamlit and used Langfuse tracing, Promptfoo evaluation, and Pytest tests in the RAG workflow.",
            ],
            "project_text_to_sql_project": [
                "Created a SQLite sales database with customer and order tables for executing generated queries.",
                "Integrated the Gemini REST API to convert natural-language questions into SQL, with an OpenAI-compatible fallback.",
                "Removed code fences from generated SQL, restricted execution to the first statement, and returned query results.",
            ],
            "project_ai_resume_jd_match_tool": [
                "Built an n8n workflow that accepts resume-file uploads and job-description input for AI-assisted analysis.",
                "Compared resume skills with job requirements and returned the matching analysis as structured output.",
            ],
        }
        for project in projects:
            bullets = rg.project_bullets(project, plan)
            self.assertEqual(bullets, expected_project_bullets[project["record_id"]])
            self.assertTrue(2 <= len(bullets) <= 3)
            self.assertTrue(all(re.match(r"^[A-Z].*[.!?]$", bullet) and len(bullet.split()) >= 8 for bullet in bullets))
        skill_groups = rg.effective_skill_groups(plan, self.profile, rg.effective_selection(plan, self.profile))
        self.assertEqual(skill_groups, [
            ("Programming", ["Python"]),
            ("Generative AI", ["Generative AI", "RAG"]),
            ("AI/LLM", ["Embeddings", "Ollama"]),
            ("Vector Database", ["ChromaDB"]),
            ("Tools / Platforms", ["Streamlit", "Langfuse", "Promptfoo", "Pytest"]),
        ])
        summary = " ".join(rg.summary_lines(plan, self.profile))
        self.assertIn("Computer Science fresher", summary)
        self.assertIn("Junior Generative AI / RAG Engineer", summary)
        for term in ("Python", "RAG", "embeddings", "ChromaDB", "Ollama", "Streamlit", "Text-to-SQL", "Gemini REST API"):
            self.assertIn(term, summary)
        self.assertIn("preprocesses student data into searchable profiles", summary)
        self.assertNotIn("verified skills in Python", summary)
        self.assertNotIn("reflectsan", summary)
        self.assertNotIn("This work reflects an", summary)

        plan["approval_checkpoint"]["resume_generation_allowed"] = True
        with tempfile.TemporaryDirectory(prefix="vectormind-rag-resume-") as temporary:
            output = Path(temporary) / "vectormind_rag_resume.docx"
            report_path = Path(temporary) / "validation.json"
            rg.generate(plan, self.profile, output)
            report = rg.validate(output, plan, self.profile, report_path)
            rendered_text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)

        jd_match = report["jd_match_validation"]
        self.assertGreater(jd_match["required"]["total"], 0)
        self.assertGreater(jd_match["preferred"]["total"], 0)
        self.assertTrue(any(item["classification"] == "required" for item in plan["jd_analysis"]["requirements"]))
        self.assertTrue(any(item["classification"] == "preferred" for item in plan["jd_analysis"]["requirements"]))
        self.assertEqual(report["final_status"], "PASS", report.get("warnings_issues"))
        self.assertTrue(report["page_count_exactly_one"])
        self.assertTrue(report["ats_validation"]["single_column"])
        self.assertTrue(report["unsupported_skills_validation"]["passed"])
        self.assertGreater(report["skills_validation"]["category_count"], 1)
        self.assertEqual(report["selected_projects"], [item["name"] for item in projects])
        self.assertEqual(report["selected_certifications"], [item["name"] for item in certifications])
        self.assertNotIn("reflectsan", rendered_text)
        for term in ("RAG", "embeddings", "ChromaDB", "Ollama", "Streamlit", "Text-to-SQL", "Gemini REST API"):
            self.assertIn(term, rendered_text)
        for bullets in expected_project_bullets.values():
            for bullet in bullets:
                self.assertIn("• " + bullet, rendered_text)
        canonical_projects = rg.effective_selection(plan, self.profile)
        self.assertEqual([item["record_id"] for item in canonical_projects], [item["record_id"] for item in projects])
        for project in canonical_projects:
            technologies = project.get("technologies") or project.get("frameworks_libraries_tools") or []
            self.assertIn("Tech Stack: " + ", ".join(technologies[:8]), rendered_text)
        skill_section = rendered_text.split("SKILLS", 1)[1].split("PROJECTS", 1)[0]
        for verified_tool in ("Langfuse", "Promptfoo", "Pytest"):
            self.assertIn(verified_tool, skill_section)
        for unsupported in ("Power BI", "Excel", "FAISS", "LangChain", "Docker", "NLP", "Prompt Engineering", "Text-to-SQL"):
            self.assertNotIn(unsupported, skill_section)


if __name__ == "__main__":
    unittest.main(verbosity=2)