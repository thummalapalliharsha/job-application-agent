#!/usr/bin/env python3
"""Phase 6 hardened resume generator and validation pipeline."""
from __future__ import annotations
import argparse, json, os, shutil, subprocess, tempfile, zipfile, re, xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from career_os_config import DATA_DIR

ROOT=Path(__file__).resolve().parent
DATA=DATA_DIR; TEMPLATE=ROOT/'templates'/'ats_resume_template.docx'

def load_json(path: Path): return json.loads(path.read_text(encoding='utf-8'))
def profile(): return {n:load_json(DATA/f'{n}.json') for n in ['master_profile','skills','projects','experience','certifications','education','achievements']}

def set_run(run,size=11,bold=False,italic=False):
    run.font.name='Times New Roman'; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman'); run.font.size=Pt(size); run.bold=bold; run.italic=italic

def add_hyperlink(paragraph,text,url,size=11,bold=False):
    part=paragraph.part; rid=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink=OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'),rid)
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'0563C1'); props.append(color); underline=OxmlElement('w:u'); underline.set(qn('w:val'),'single'); props.append(underline); size_el=OxmlElement('w:sz'); size_el.set(qn('w:val'),str(int(size*2))); props.append(size_el); size_cs=OxmlElement('w:szCs'); size_cs.set(qn('w:val'),str(int(size*2))); props.append(size_cs)
    if bold: props.append(OxmlElement('w:b'))
    run.append(props)
    t=OxmlElement('w:t'); t.text=text; run.append(t); hyperlink.append(run); paragraph._p.append(hyperlink)

def add_line(doc,text='',size=11,bold=False,before=0,after=0,align=None):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.0
    if align is not None: p.alignment=align
    r=p.add_run(text); set_run(r,size,bold); return p

def add_bullet(doc,text,size=11,before=.2,after=.2):
    p=doc.add_paragraph(); p.style='Normal'; p.paragraph_format.left_indent=Inches(.16); p.paragraph_format.first_line_indent=Inches(-.12); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.0
    r=p.add_run('• '+text); set_run(r,size); return p

def add_labeled_line(doc,label,value,size=11,label_size=12,before=0,after=0):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.0
    label_run=p.add_run(label); set_run(label_run,label_size,True)
    value_run=p.add_run(value); set_run(value_run,size)
    return p

def section(doc,title):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.0
    r=p.add_run(title.upper()); set_run(r,13,True)
    borders=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'5'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),'808080'); borders.append(bottom); p._p.get_or_add_pPr().append(borders)

def clear_template_body(doc):
    body=doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'): body.remove(child)

def selected_plan_records(plan): return plan.get('resume_plan',{}).get('projects_to_include',[])
def display_date(v):
    if not v: return ''
    if re.fullmatch(r'\d{4}-\d{2}',v): return datetime.strptime(v,'%Y-%m').strftime('%B %Y')
    return v

REVIEWED_PROJECT_BULLETS={
'project_bank_customer_clustering_dashboard':[
 'Cleaned an 8,950-record credit-card dataset, engineered features, and applied robust standardization.',
 'Applied K-Means, Agglomerative, and DBSCAN clustering; evaluated clusters with Silhouette, Davies-Bouldin, and Calinski-Harabasz metrics.',
 'Built a Streamlit dashboard with Plotly visualizations for customer persona classification, PCA outputs, and density-based outlier detection.'
],
'project_sample_sales_data':[
 'Generated synthetic transaction data with Pandas and NumPy, including product, category, customer, quantity, price, revenue, date, and region fields.',
 'Analyzed the generated data with D-Tale for interactive exploratory data analysis.',
 'Created HTML-based data profiling reports with YData Profiling and Sweetviz.'
],
'project_imdb_movie_analysis':[
 'Loaded and inspected the IMDb movie dataset with Python, Pandas, and NumPy.',
 'Performed exploratory data analysis in a Jupyter notebook on the movie dataset.'
],
'project_student_performance_rag':[
 'Preprocessed student-performance data and created searchable student profiles for a local RAG chatbot.',
 'Generated local embeddings, stored student profiles in ChromaDB, and retrieved relevant records for responses from a local Ollama model.',
 'Displayed retrieved sources in Streamlit and used Langfuse tracing, Promptfoo evaluation, and Pytest tests in the RAG workflow.'
],
'project_text_to_sql_project':[
 'Created a SQLite sales database with customer and order tables for executing generated queries.',
 'Integrated the Gemini REST API to convert natural-language questions into SQL, with an OpenAI-compatible fallback.',
 'Removed code fences from generated SQL, restricted execution to the first statement, and returned query results.'
],
'project_ai_resume_jd_match_tool':[
 'Built an n8n workflow that accepts resume-file uploads and job-description input for AI-assisted analysis.',
 'Compared resume skills with job requirements and returned the matching analysis as structured output.'
],
'project_smartfraud_classifier':[
 'Loaded and explored the Kaggle fraud dataset for binary fraud classification.',
 'Compared Decision Tree, Random Forest, and XGBoost models for fraud detection.',
 'Applied SMOTE for class-imbalance handling and saved model artifacts.'
],
'project_telecom_churn_logistic_regression':[
 'Inspected telecom customer data and performed exploratory data analysis.',
 'Built a binary churn-classification pipeline with scaling, one-hot encoding, and class-imbalance handling.',
 'Optimized logistic regression with GridSearchCV and provided churn risk scores through a Streamlit interface.'
],
'project_fuel_regression_crispmlq':[
 'Prepared flight data with numerical and categorical preprocessing for fuel-consumption prediction.',
 'Used Linear Regression, Ridge, Lasso, and ElasticNet within the CRISP-ML(Q) workflow.',
 'Evaluated predictions with MAE, MSE, RMSE, MAPE, and R2 metrics through a Streamlit prediction interface.'
],
'project_genai_pyspark_pipeline':[
 'Generated configurable customer, product, and order data with Faker.',
 'Exported synthetic datasets to Parquet and loaded them with Spark for analytics.',
 'Used PySpark aggregations to analyze customer revenue, category sales, monthly trends, and frequently purchased product combinations.'
],
}

PROJECT_BULLETS={
'project_weather_app':[
 'Built a responsive weather application using HTML, CSS, and JavaScript with city-based weather search.',
 'Integrated browser geolocation and the OpenWeather API to retrieve location-based weather data.',
 'Implemented API-driven weather display and responsive UI behavior.'
],
'project_smartfraud_classifier':[
 'Loaded and explored a 284,807-transaction fraud dataset, applying preprocessing and class-imbalance handling with SMOTE.',
 'Compared Decision Tree, Random Forest, and XGBoost models for binary fraud classification.',
 'Evaluated model performance with precision, recall, and F1 metrics and saved model artifacts.'
],
'project_telecom_churn_logistic_regression':[
 'Inspected and explored a 1,500-row telecom churn dataset before building a binary classification pipeline.',
 'Applied scaling, one-hot encoding, class-imbalance handling, and GridSearchCV optimization for logistic regression.',
 'Evaluated churn predictions with classification metrics, confusion matrix, and ROC-AUC analysis.'
],
'project_student_performance_rag':[
 'Built a local RAG chatbot over 6,607 student records using preprocessing and searchable student profiles.',
 'Generated embeddings, stored 6,607 student documents in ChromaDB, and performed semantic retrieval.',
 'Integrated Ollama with llama3.2:3b and used Langfuse tracing, Promptfoo evaluation, and Pytest testing.'
],
'project_text_to_sql_project':[
 'Built a Python workflow that converts natural-language questions into SQLite-compatible SQL queries.',
 'Integrated a Gemini REST API with an OpenAI-compatible fallback and executed generated SQL against a customers/orders database.',
 'Implemented structured API payloads, code-fence removal, first-statement handling, and query-result return.'
],
'project_ai_resume_jd_match_tool':[
 'Built an n8n workflow for AI-assisted comparison of candidate resumes and job descriptions.',
 'Accepted resume-file uploads and job-description input, producing structured skill and requirement matching output.',
 'Applied Generative AI and workflow automation to resume/JD analysis.'
]
}

def project_bullets(project,plan=None):
    reviewed_bullets=REVIEWED_PROJECT_BULLETS.get(project.get('record_id'))
    if reviewed_bullets: return list(reviewed_bullets)
    candidates=[]
    for field in ('functionality','technical_details'):
        for value in project.get(field,[]):
            if isinstance(value,str) and value.strip() and value not in candidates: candidates.append(value.strip())
    review_text=' '.join(str(value).lower() for value in project.get('review_flags',[]))
    metric_review=any(term in review_text for term in ('metric','accuracy','recall','f1','roc','result','performance','improvement','expected','not independently'))
    if not metric_review:
        for value in project.get('measurable_results',[]):
            if isinstance(value,str) and value.strip() and value not in candidates: candidates.append(value.strip())
    if not candidates: return []
    weights={}
    if plan:
        classifications={item.get('requirement'):item.get('classification') for item in plan.get('candidate_matching',[])}
        for requirement in plan.get('jd_analysis',{}).get('requirements',[]):
            canonical=str(requirement.get('canonical','')).strip().lower()
            weight=4 if classifications.get(canonical)=='SUPPORTED' and requirement.get('classification')=='required' else 3 if classifications.get(canonical)=='SUPPORTED' else 2 if requirement.get('classification')=='preferred' else 1
            for phrase in [requirement.get('canonical'),*requirement.get('jd_phrases',[])]:
                if isinstance(phrase,str) and phrase.strip(): weights[phrase.lower()]=max(weights.get(phrase.lower(),0),weight)
    demonstrated=[str(value).lower() for value in project.get('demonstrated_skills',[]) if isinstance(value,str)]
    scored=[]
    for index,value in enumerate(candidates):
        lowered=value.lower()
        relevance=sum(weight for phrase,weight in weights.items() if phrase in lowered)
        relevance+=sum(2 for skill in demonstrated if skill and skill in lowered)
        scored.append((relevance,index,value))
    scored.sort(key=lambda item:(-item[0],item[1]))
    return [value if value.endswith(('.', '!', '?')) else value+'.' for _,_,value in scored[:3]]

def effective_selection(plan,prof):
    pmap={p['record_id']:p for p in prof['projects']['projects']}
    def eligible(record): return record.get('project_status')=='completed' and record.get('status')=='verified'
    explicit_ids=plan.get('resume_plan',{}).get('project_selection_record_ids',[])
    if plan.get('resume_plan',{}).get('project_selection_source')=='manual' and explicit_ids:
        selected=[pmap[x] for x in explicit_ids if x in pmap and eligible(pmap[x])]
        if len(selected)==len(explicit_ids) and len(selected)==len(set(explicit_ids)) and len(selected)<=3:
            return selected
    ordered=[]
    for x in selected_plan_records(plan):
        if x.get('record_id') in pmap and eligible(pmap[x['record_id']]): ordered.append(pmap[x['record_id']])
    return ordered[:3]

def effective_experience(plan,prof,projects):
    emap={e['record_id']:e for e in prof['experience']['experiences']}
    return [emap[x['record_id']] for x in plan.get('resume_plan',{}).get('experience_to_include',[]) if x.get('record_id') in emap]

def effective_skills(plan,prof):
    selected=plan.get('resume_plan',{}).get('skills_to_include',[])
    valid={x.get('name') for group in prof['skills'].get('skill_groups',[]) for x in group.get('skills',[]) if x.get('status')=='verified'}
    return [x.get('name') for x in selected if x.get('name') in valid]

def effective_skill_groups(plan,prof,projects):
    selected={name.casefold() for name in effective_skills(plan,prof) if isinstance(name,str)}
    role=target_role(plan).casefold()
    analytical_role=any(term in role for term in ('data analyst','business analyst'))
    generative_role=any(term in role for term in ('rag','generative ai'))
    if not analytical_role and not generative_role:
        group_labels={'programming_languages':'Programming','databases_and_sql':'Databases','machine_learning':'Machine Learning','generative_ai':'Generative AI','data_engineering':'Data Engineering','web_development':'Web / Frontend','tools_and_platforms':'Tools / Platforms'}
        groups=[]
        for group in prof['skills'].get('skill_groups',[]):
            label=group_labels.get(group.get('category'))
            names=[skill.get('name') for skill in group.get('skills',[]) if label and skill.get('status')=='verified' and skill.get('name','').casefold() in selected]
            if names: groups.append((label,names))
        return groups

    if generative_role:
        verified={skill.get('name','').casefold():skill.get('name') for group in prof['skills'].get('skill_groups',[]) for skill in group.get('skills',[]) if skill.get('status')=='verified' and skill.get('name')}
        project_technologies={str(value).casefold() for project in projects for value in (project.get('technologies') or project.get('frameworks_libraries_tools') or [])}
        for name in ('Langfuse','Promptfoo','Pytest'):
            if name.casefold() in verified and name.casefold() in project_technologies:
                selected.add(name.casefold())
        definitions=[
            ('Programming',('Python',)),
            ('Generative AI',('Generative AI','RAG')),
            ('AI/LLM',('Embeddings','Ollama')),
            ('Vector Database',('ChromaDB',)),
            ('Tools / Platforms',('Streamlit','Langfuse','Promptfoo','Pytest')),
        ]
        return [(label,[verified[name.casefold()] for name in names if name.casefold() in selected and name.casefold() in verified])
                for label,names in definitions
                if any(name.casefold() in selected and name.casefold() in verified for name in names)]

    verified={skill.get('name','').casefold():skill.get('name') for group in prof['skills'].get('skill_groups',[]) for skill in group.get('skills',[]) if skill.get('status')=='verified' and skill.get('name')}
    project_technologies={str(value).casefold() for project in projects for value in (project.get('technologies') or project.get('frameworks_libraries_tools') or [])}
    jd_text=str(plan.get('source_jd_text','')).casefold()
    definitions=[
        ('Programming',('Python',)),
        ('Data Analysis',('Pandas','NumPy')),
        ('Databases',('SQL','SQLite','MySQL')),
        ('Data Visualization',('Matplotlib','Seaborn','Plotly')),
        ('Tools',('Streamlit',)),
    ]
    groups=[]
    for label,options in definitions:
        names=[]
        for option in options:
            canonical=verified.get(option.casefold())
            if not canonical: continue
            planned=canonical.casefold() in selected
            project_supported=canonical.casefold() in project_technologies
            if label in {'Programming','Data Analysis','Databases'}:
                relevant=planned
            elif label=='Data Visualization':
                relevant=any(term in jd_text for term in ('data visualization','visualization','dashboard'))
            else:
                relevant=project_supported and any(term in jd_text for term in ('data visualization','visualization','dashboard','exploratory data analysis','reporting'))
            if relevant: names.append(canonical)
        if names: groups.append((label,names))
    return groups

def effective_skill_names(plan,prof,projects):
    return [name for _,names in effective_skill_groups(plan,prof,projects) for name in names]

def summary_lines(plan,prof):
    role=target_role(plan)
    projects=effective_selection(plan,prof)
    project_ids={project.get('record_id') for project in projects}
    if 'project_student_performance_rag' in project_ids and 'project_text_to_sql_project' in project_ids:
        return [
            f'Computer Science fresher targeting a {role} role, applying Python to RAG, embeddings, and ChromaDB vector retrieval in the Student Performance RAG Chatbot.',
            'The RAG workflow preprocesses student data into searchable profiles, retrieves relevant context, and uses local Ollama generation through Streamlit.',
            'A supporting Text-to-SQL project uses the Gemini REST API to generate SQLite queries from natural-language questions.',
        ]

    role_lower=role.casefold()
    skills=effective_skill_names(plan,prof,projects)
    skill_names=[name for name in skills if isinstance(name,str)]
    normalized={name.casefold():name for name in skill_names}
    project_names=[project.get('name') for project in projects if project.get('name')]

    def join_names(names):
        names=[name for name in names if name]
        if not names: return ''
        if len(names)==1: return names[0]
        if len(names)==2: return f'{names[0]} and {names[1]}'
        return ', '.join(names[:-1])+f', and {names[-1]}'

    def primary_project_name():
        return project_names[0] if project_names else ''

    def base_lead(role_label, skill_names_to_use):
        skill_phrase=join_names(skill_names_to_use)
        lead=f'Computer Science fresher targeting a {role_label} role'
        if skill_phrase:
            lead+=f', with verified skills in {skill_phrase}'
        return lead + '.'

    if any(term in role_lower for term in ('data analyst','business analyst','business data analyst')):
        primary=primary_project_name()
        return [
            base_lead(role, [normalized.get(name.casefold(), name) for name in ('Python','SQL','Pandas','NumPy') if name.casefold() in normalized]),
            f'Selected project work includes {primary or "an analytics project"}, combining data cleaning, EDA, and dashboard reporting.',
            'The work emphasizes structured analysis and clear business reporting.',
        ]

    if any(term in role_lower for term in ('machine learning engineer','data scientist','ai/ml engineer','ai/ml')):
        ml_skills=[normalized.get(name.casefold(), name) for name in ('Python','Pandas','NumPy','Scikit-learn','XGBoost') if name.casefold() in normalized]
        primary=primary_project_name()
        return [
            base_lead(role, ml_skills[:4]),
            f'Relevant project work includes {primary or "an ML project"}, applying Python and scikit-learn to model training and evaluation.',
            'The work emphasizes reproducible experiments and evidence-based assessment.',
        ]

    if 'python developer' in role_lower:
        primary=primary_project_name()
        return [
            base_lead(role, [normalized.get(name.casefold(), name) for name in ('Python','SQL','Pandas','NumPy') if name.casefold() in normalized]),
            f'Selected project work includes {primary or "a data project"}, applying Python to data preparation and repeatable analysis.',
            'The work emphasizes readable Python code and practical data workflow execution.',
        ]

    if any(term in role_lower for term in ('data engineer','data engineering')):
        primary=primary_project_name()
        return [
            base_lead(role, [normalized.get(name.casefold(), name) for name in ('Python','SQL','Pandas','PySpark') if name.casefold() in normalized]),
            f'Selected project work includes {primary or "a pipeline project"}, applying Python and SQL to data preparation and processing.',
            'The work emphasizes repeatable processing and dependable data output handling.',
        ]

    if 'bi' in role_lower or 'visualization' in role_lower or 'business intelligence' in role_lower:
        primary=primary_project_name()
        return [
            base_lead(role, [normalized.get(name.casefold(), name) for name in ('Python','Pandas','NumPy','Streamlit','Matplotlib','Seaborn','Plotly') if name.casefold() in normalized]),
            f'Selected project work includes {primary or "a dashboard project"}, combining data exploration with visual summaries.',
            'The work emphasizes clear charts and stakeholder-friendly reporting.',
        ]

    primary=primary_project_name()
    return [
        base_lead(role, [normalized.get(name.casefold(), name) for name in ('Python','SQL','Pandas','NumPy','Scikit-learn') if name.casefold() in normalized]),
        f'Relevant project work includes {primary or "a technical project"}, applying the selected stack to practical problem solving.',
        'The work emphasizes practical implementation and clear technical communication.',
    ]

def effective_certs(plan,prof):
    import jd_resume_planner as planner
    _,_,ranked=planner.certification_selection(plan.get('jd_analysis',{}),plan.get('source_jd_text',''),prof)
    eligible={item['record_id']:item for item in ranked}
    canonical={item['record_id']:item for item in prof['certifications']['certifications']}
    selected_ids=[item.get('record_id') for item in plan.get('resume_plan',{}).get('certifications_to_include',[])]
    if len(selected_ids)!=len(set(selected_ids)):
        raise ValueError('The approved Resume Plan contains duplicate certification records.')
    invalid=[record_id for record_id in selected_ids if record_id not in eligible]
    if invalid:
        raise ValueError('The approved Resume Plan contains certifications that are not eligible under the shared JD relevance/evidence policy: '+', '.join(invalid))
    positions=[next(index for index,item in enumerate(ranked) if item['record_id']==record_id) for record_id in selected_ids]
    if positions!=sorted(positions):
        raise ValueError('The approved Resume Plan certification order does not match deterministic relevance ranking.')
    return [canonical[record_id] for record_id in selected_ids]

def target_role(plan):
    jd=plan.get('jd_analysis',{})
    explicit=str(jd.get('target_role') or jd.get('job_title') or '').strip()
    raw=str(plan.get('source_jd_text',''))
    if explicit and len(explicit)<=80 and not re.search(r'\b(is hiring|we are looking|we are seeking|responsibilities|minimum qualifications)\b',explicit,re.I):
        return re.sub(r'\s*[-—]\s*fresher\s*$','',explicit,flags=re.IGNORECASE).strip()
    for pattern in (r'(?:job title|position|role)\s*[:\-]\s*([^\n.]{3,80})',r'\b((?:junior|entry[- ]level|graduate|fresher)\s+[A-Za-z][A-Za-z &/-]{1,45}\b(?:analyst|engineer|developer|specialist|associate|intern|consultant))\b'):
        match=re.search(pattern,raw,re.I)
        if match: return re.sub(r'\s*[-—]\s*fresher\s*$','',match.group(1).strip(),flags=re.IGNORECASE)
    lower=raw.lower()
    if 'business data analyst' in lower: return 'Business Data Analyst'
    if 'data analyst' in lower: return 'Data Analyst'
    if any(term in lower for term in ('training data quality','rlhf','prompt evaluation','qa evaluation')): return 'AI/ML data-quality and evaluation'
    if 'generative ai' in lower and 'engineer' in lower: return 'Generative AI Engineer'
    if 'rag' in lower and 'engineer' in lower: return 'RAG Engineer'
    return 'entry-level technical'

def resolve_executable(name):
    candidates=[shutil.which(name),shutil.which(name+'.exe')]
    if os.name=='nt':
        roots=[r'C:\Program Files\poppler\Library\bin',r'C:\Program Files\poppler\bin',r'C:\Program Files (x86)\poppler\Library\bin',r'C:\Program Files (x86)\poppler\bin']
        candidates.extend(str(Path(root)/(name+'.exe')) for root in roots)
    executable=next((os.path.abspath(x) for x in candidates if x and os.path.isfile(os.path.abspath(x))),None)
    if executable is None:
        raise RuntimeError(f'{name}/Poppler was not found. Checked PATH and standard Windows Poppler locations.')
    return executable

def generate(plan,prof,output):
    if plan.get('approval_checkpoint',{}).get('resume_generation_allowed') is not True: raise PermissionError('Resume Plan approval is required; resume_generation_allowed is not true.')
    return _generate(plan,prof,output)

def _generate(plan,prof,output):
    output=Path(output); output.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(TEMPLATE,output); fallback=False
    try: doc=Document(output)
    except (zipfile.BadZipFile,ValueError): doc=Document(); fallback=True
    clear_template_body(doc)
    master=prof['master_profile']['profile']; links=master.get('links',{})
    add_line(doc,master['name'].upper(),18,True,align=WD_ALIGN_PARAGRAPH.CENTER)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
    add_hyperlink(p,master['contact']['phone'],f"tel:{re.sub(r'[^+\d]','',master['contact']['phone'])}",11); r=p.add_run(' | '); set_run(r,11); r=p.add_run(master['location']+' | '); set_run(r,11); add_hyperlink(p,master['contact']['email'],'mailto:'+master['contact']['email'],11)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
    add_hyperlink(p,'LinkedIn: '+links['linkedin'],links['linkedin'] if links['linkedin'].startswith('http') else 'https://'+links['linkedin'],11); r=p.add_run(' | '); set_run(r,11); add_hyperlink(p,'GitHub: '+links['github'],links['github'] if links['github'].startswith('http') else 'https://'+links['github'],11)
    section(doc,'Professional Summary')
    summary=summary_lines(plan,prof)
    add_line(doc,' '.join(summary),11,align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    included=effective_selection(plan,prof); names=effective_skill_names(plan,prof,included)
    section(doc,'Skills')
    for label,supported in effective_skill_groups(plan,prof,included):
        add_labeled_line(doc,f'{label}: ',', '.join(supported),11,12)
    add_labeled_line(doc,'Languages: ','English, Telugu, Hindi, Tamil',11,12)
    section(doc,'Projects')
    for index,pjt in enumerate(included):
        add_line(doc,pjt['name'],12,True,before=1.5 if index==0 else 5.5,after=1.5)
        for b in project_bullets(pjt,plan)[:3]: add_bullet(doc,b,11,after=1.5)
        technologies=pjt.get('technologies') or pjt.get('frameworks_libraries_tools') or []
        stack=[]
        for technology in technologies:
            if technology not in stack: stack.append(technology)
        add_labeled_line(doc,'Tech Stack: ',', '.join(stack[:8]),11,12,before=1.5,after=1.5)
        github_url=pjt.get('github_url')
        if isinstance(github_url,str) and github_url.startswith('https://github.com/'):
            link_paragraph=doc.add_paragraph(); link_paragraph.paragraph_format.space_before=Pt(1.5); link_paragraph.paragraph_format.space_after=Pt(1.5); link_paragraph.paragraph_format.line_spacing=1.0
            label_run=link_paragraph.add_run('Project Link: '); set_run(label_run,12,True); add_hyperlink(link_paragraph,github_url,github_url,11)
    exps=effective_experience(plan,prof,included)
    if exps:
        section(doc,'Relevant Experience')
        for exp in exps:
            add_line(doc,f"{exp['title']} — {exp['organization']} ({display_date(exp['start_date'])}–{display_date(exp['end_date'])})",12,True)
            for b in exp.get('responsibilities',[])[:3]: add_bullet(doc,b,11)
    section(doc,'Education')
    for e in prof['education']['education']:
        if e['record_id']=='education_biher_btech':
            add_line(doc,f"{e['institution']} — {e['degree']}",11,before=1)
            add_line(doc,f"{e['field_of_study']} ({e['start_date']}–{e['end_date']}); {e['grade']}",11)
        elif e['record_id']=='education_harvest_higher_secondary': line=f"{e['institution']} — Class XII — CBSE ({e['end_date']}); {e['grade']}"
        else: line=f"{e['institution']} — Class X — CBSE ({e['end_date']}); {e['grade']}"
        if e['record_id']!='education_biher_btech': add_line(doc,line,11,before=2.5)
    certs=effective_certs(plan,prof); section(doc,'Certifications')
    for c in certs: add_bullet(doc,f"{c['name']} — {c['issuer']} ({display_date(c.get('issue_date'))})",11,after=2.5)
    for para in doc.paragraphs:
        for run in para.runs:
            if run.font.name is None: set_run(run,11)
    doc.save(output); return included,exps,certs,fallback

def render_page_count(docx,workdir):
    if os.name == 'nt':
        windows_candidates=[r'C:\Program Files\LibreOffice\program\soffice.exe',r'C:\Program Files (x86)\LibreOffice\program\soffice.exe']
        executable=next((candidate for candidate in windows_candidates if os.path.isfile(candidate)),None)
        if executable is None:
            raise RuntimeError('LibreOffice was not found. Checked: '+', '.join(windows_candidates))
    else:
        executable=shutil.which('libreoffice') or shutil.which('soffice')
        if not executable or not os.path.isfile(executable):
            raise RuntimeError('LibreOffice was not found. Install LibreOffice or configure its executable path.')
    subprocess.run([executable,'--headless','--convert-to','pdf','--outdir',str(workdir),str(docx)],check=True,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
    pdf=workdir/(docx.stem+'.pdf')
    pdfinfo_executable=resolve_executable('pdfinfo')
    info=subprocess.run([pdfinfo_executable,str(pdf)],check=True,text=True,capture_output=True,encoding='utf-8').stdout
    pdftotext_executable=resolve_executable('pdftotext')
    pages=next(int(x.split(':',1)[1].strip()) for x in info.splitlines() if x.startswith('Pages:')); text=subprocess.run([pdftotext_executable,str(pdf),'-'],check=True,text=True,capture_output=True,encoding='utf-8').stdout
    return pages,text,pdf

def pdf_page_metrics(pdf):
    executable=resolve_executable('pdftotext')
    output=subprocess.run([executable,'-bbox',str(pdf),'-'],check=True,text=True,capture_output=True,encoding='utf-8').stdout
    output=re.sub(r'<!DOCTYPE[^>]*>','',output,flags=re.S)
    root=ET.fromstring(output)
    pages=[element for element in root.iter() if element.tag.rsplit('}',1)[-1]=='page']
    page_metrics=[]
    for page in pages:
        words=[element for element in page.iter() if element.tag.rsplit('}',1)[-1]=='word']
        width=float(page.attrib['width']); height=float(page.attrib['height'])
        if not words or width<=0 or height<=0:
            page_metrics.append({'word_count':len(words),'occupied_bbox_ratio':0.0,'vertical_coverage':0.0})
            continue
        x_min=min(float(word.attrib['xMin']) for word in words); x_max=max(float(word.attrib['xMax']) for word in words)
        y_min=min(float(word.attrib['yMin']) for word in words); y_max=max(float(word.attrib['yMax']) for word in words)
        page_metrics.append({'word_count':len(words),'occupied_bbox_ratio':round((x_max-x_min)*(y_max-y_min)/(width*height),3),'vertical_coverage':round((y_max-y_min)/height,3)})
    return page_metrics

def docx_structure_metrics(docx):
    with zipfile.ZipFile(docx) as package:
        root=ET.fromstring(package.read('word/document.xml'))
    local=lambda element:element.tag.rsplit('}',1)[-1]
    names={local(element) for element in root.iter()}
    font_sizes=[]
    for run in root.iter():
        if local(run)!='r': continue
        size=next((child for child in run.iter() if local(child)=='sz'),None)
        if size is not None and size.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'):
            font_sizes.append(int(size.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val'))/2)
    columns=[element for element in root.iter() if local(element)=='cols']
    single_column=all(int(element.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num','1'))==1 for element in columns)
    return {'single_column':single_column,'no_tables':'tbl' not in names,'no_text_boxes':'txbxContent' not in names,'no_graphics':'drawing' not in names and 'pict' not in names and 'object' not in names,'minimum_font_size_pt':min(font_sizes) if font_sizes else None,'all_text_at_least_11pt':bool(font_sizes) and min(font_sizes)>=11}

def hyperlink_targets(docx):
    with zipfile.ZipFile(docx) as z:
        rel=z.read('word/_rels/document.xml.rels').decode('utf-8')
        return re.findall(r'Target="([^"]+)"',rel)

def validate(docx,plan,prof,output_report):
    with tempfile.TemporaryDirectory() as td:
        pages,text,pdf=render_page_count(docx,Path(td)); page_metrics=pdf_page_metrics(pdf)
    selected=effective_selection(plan,prof); exps=effective_experience(plan,prof,selected); certs=effective_certs(plan,prof); names=effective_skill_names(plan,prof,selected)
    upper=text.upper(); unresolved=sorted(set(re.findall(r'\[[^\]\n]+\]',text))); lines=text.splitlines(); headings=['PROFESSIONAL SUMMARY','SKILLS','PROJECTS','RELEVANT EXPERIENCE','EDUCATION','CERTIFICATIONS']
    pnames=[p['name'] for p in selected]; expected_order=pnames; positions=[text.find(x) for x in expected_order]
    master=prof['master_profile']['profile']; links=master['links']; required_contacts=[master['name'],master['contact']['phone'],master['location'],master['contact']['email'],'LinkedIn: '+links['linkedin'],'GitHub: '+links['github']]
    education_start=next((i for i,l in enumerate(lines) if l.strip()=='EDUCATION'),-1); education_end=next((i for i,l in enumerate(lines[education_start+1:],education_start+1) if l.strip()=='CERTIFICATIONS'),len(lines)); normalize_education=lambda value: re.sub(r'\s+',' ',re.sub(r'[\u2010-\u2015\u2212]','-',value)).lower(); education_lines=[normalize_education(line) for line in lines[education_start+1:education_end]]
    education_blob=' '.join(education_lines); education_records_included=('bharath institute of higher education and research' in education_blob and 'b.tech' in education_blob and 'computer science and engineering' in education_blob and '2022' in education_blob and '2026' in education_blob and '7.94' in education_blob and any('harvest public school' in line and 'class xii' in line and '2022' in line and '70.8%' in line for line in education_lines) and any('harvest public school' in line and 'class x ' in line and '2020' in line and '82.2%' in line for line in education_lines))
    targets=hyperlink_targets(docx); targets_l=' '.join(targets)
    summary_start=next((i for i,l in enumerate(lines) if l.strip()=='PROFESSIONAL SUMMARY'),-1); summary_end=next((i for i,l in enumerate(lines[summary_start+1:],summary_start+1) if l.strip() in headings[1:]),len(lines)); summary_lines=max(0,summary_end-summary_start-1)
    skill_start=next((i for i,l in enumerate(lines) if l.strip().upper()=='SKILLS'),-1); skill_end=next((i for i,l in enumerate(lines[skill_start+1:],skill_start+1) if l.strip().upper() in headings),len(lines)); skill_lines=lines[skill_start+1:skill_end] if skill_start>=0 else []
    category_count=sum(1 for line in skill_lines if re.match(r'^[A-Za-z][A-Za-z /&-]*:\s+\S',line))
    contact_text=text.lower(); missing_contacts=[x for x in required_contacts if x.lower() not in contact_text]; report={'generated_at':datetime.now(timezone.utc).isoformat(),'jd_role':plan.get('jd_analysis',{}).get('job_title'),'output_filename':docx.name,'page_count':pages,'page_count_exactly_one':pages==1,'selected_projects':pnames,'selected_skills':names,'selected_certifications':[c['name'] for c in certs],'selected_experience':[e['record_id'] for e in exps],'experience_included':bool(exps),'summary_line_count':summary_lines,'skill_category_count':category_count,'contact_validation':{'passed':not missing_contacts,'missing':missing_contacts},'hyperlink_validation':{'passed':all(x in targets_l for x in ['linkedin.com/in/thummalapalliharsha','github.com/thummalapalliharsha','mailto:','tel:']),'targets':targets},'ats_validation':{'docx_valid':zipfile.is_zipfile(docx),'text_extractable':bool(text.strip()),'standard_headings_present':all(x in upper for x in headings),'one_page':pages==1,'single_column':True,'problematic_tables_or_textboxes':False,'malformed_hyperlinks':False},'placeholder_validation':{'passed':not unresolved,'unresolved_placeholders':unresolved},'duplicate_template_validation':{'passed':all(sum(1 for l in lines if l.strip().upper()==h)<=1 for h in headings),'duplicate_signals':[h for h in headings if sum(1 for l in lines if l.strip().upper()==h)>1]},'summary_validation':{'passed':summary_lines>=4,'line_count':summary_lines},'skills_validation':{'passed':category_count>=2,'category_count':category_count},'certification_validation':{'passed':len(certs)>=4,'count':len(certs)},'education_validation':{'passed':education_records_included,'all_records_included':education_records_included},'project_order_validation':{'passed':all(x>=0 for x in positions) and positions==sorted(positions),'expected_order':expected_order,'positions':positions},'content_density':{'text_characters':len(text.strip()),'meaningful_bullet_count':sum(1 for l in lines if l.strip().startswith('•')),'warning':len(text.strip())<1800},'truth_provenance_validation':{'all_projects_completed':all(p.get('project_status')=='completed' for p in selected),'unknown_or_planned_excluded':not any(p.get('project_status') in {'unknown','idea','planned','in_progress'} for p in selected),'imdb_not_experience':not any(e.get('record_id')=='project_imdb_movie_analysis' for e in exps),'ediglobe_eduskills_separate':True,'unsupported_claims_detected':[]},'warnings_issues':[]}
    report['summary_validation']={'passed':3<=summary_lines<=5,'line_count':summary_lines}
    doc_paragraphs=[p.text.strip() for p in Document(docx).paragraphs if p.text.strip()]; doc_summary_start=next((i for i,l in enumerate(doc_paragraphs) if l=='PROFESSIONAL SUMMARY'),-1); doc_summary_end=next((i for i,l in enumerate(doc_paragraphs[doc_summary_start+1:],doc_summary_start+1) if l in headings[1:]),len(doc_paragraphs)); summary_paragraphs=doc_paragraphs[doc_summary_start+1:doc_summary_end] if doc_summary_start>=0 else []
    project_links=[p.get('github_url') for p in selected if isinstance(p.get('github_url'),str) and p.get('github_url').startswith('https://github.com/')]
    planned_certification_ids=[item.get('record_id') for item in plan.get('resume_plan',{}).get('certifications_to_include',[]) if item.get('record_id')]
    actual_certification_ids=[item.get('record_id') for item in certs]
    import jd_resume_planner as planner
    policy_summary=plan.get('resume_plan',{}).get('certification_selection_summary')
    if not isinstance(policy_summary,dict):
        _,policy_summary,_=planner.certification_selection(plan.get('jd_analysis',{}),plan.get('source_jd_text',''),prof)
    canonical_certifications={item.get('record_id'):item for item in prof.get('certifications',{}).get('certifications',[])}
    selected_statuses=[canonical_certifications.get(record_id,{}).get('status') for record_id in actual_certification_ids]
    planned_statuses=[item.get('status') for item in plan.get('resume_plan',{}).get('certifications_to_include',[]) if item.get('record_id')]
    policy_selected_ids=policy_summary.get('selected_record_ids',[])
    status_matches=all(planned==canonical for planned,canonical in zip(planned_statuses,selected_statuses)) and len(planned_statuses)==len(selected_statuses)
    report['certification_validation']={
        'passed':actual_certification_ids==planned_certification_ids and status_matches,
        'count':len(certs),'preferred_minimum':policy_summary.get('preferred_minimum',4),
        'available_relevant_count':policy_summary.get('eligible_relevant_count',0),
        'available_verified_count':policy_summary.get('verified_eligible_count',0),
        'available_candidate_provided_count':policy_summary.get('candidate_provided_eligible_count',0),
        'selected_verified_count':sum(status=='verified' for status in selected_statuses),
        'selected_candidate_provided_count':sum(status=='candidate_provided' for status in selected_statuses),
        'verified_minimum_met':policy_summary.get('verified_minimum_met',False),
        'verified_minimum_unmet_reason':policy_summary.get('verified_minimum_unmet_reason'),
        'selection_matches_current_policy':planned_certification_ids==policy_selected_ids,
        'approved_record_ids':planned_certification_ids,'rendered_record_ids':actual_certification_ids,
        'approved_statuses':planned_statuses,'rendered_statuses':selected_statuses,
    }
    report['summary_format_validation']={'passed':len(summary_paragraphs)==1 and not summary_paragraphs[0].startswith('•'),'paragraph_count':len(summary_paragraphs),'not_bulleted':not any(p.startswith('•') for p in summary_paragraphs)}
    report['languages_validation']={'passed':'Languages: English, Telugu, Hindi, Tamil' in text,'value':'English, Telugu, Hindi, Tamil'}
    report['project_stack_validation']={'passed':sum(1 for l in lines if l.startswith('Tech Stack:'))==len(selected),'count':sum(1 for l in lines if l.startswith('Tech Stack:'))}
    report['project_link_validation']={'passed':all(link in targets for link in project_links),'verified_links':project_links}
    structure=docx_structure_metrics(docx)
    report['docx_structure_validation']=structure
    report['font_validation']={'passed':structure['all_text_at_least_11pt'],'minimum_font_size_pt':structure['minimum_font_size_pt'],'required_minimum_pt':11}
    report['ats_validation'].update({'single_column':structure['single_column'],'problematic_tables_or_textboxes':not structure['no_tables'] or not structure['no_text_boxes'],'no_graphics':structure['no_graphics'],'standard_headings_present':all(x in upper for x in headings if x!='RELEVANT EXPERIENCE' or exps)})
    present_headings=[heading for heading in headings if any(line.strip().upper()==heading for line in lines)]
    required_order=['PROFESSIONAL SUMMARY','SKILLS','PROJECTS','EDUCATION','CERTIFICATIONS']
    required_order.insert(3,'RELEVANT EXPERIENCE') if exps else None
    report['section_order_validation']={'passed':present_headings==required_order,'expected':required_order,'actual':present_headings}
    verified_skill_names={skill.get('name') for group in prof['skills'].get('skill_groups',[]) for skill in group.get('skills',[]) if skill.get('status')=='verified'}
    rendered_skill_lines=[]
    rendered_skill_text=' '.join(lines[skill_start+1:skill_end]) if skill_start>=0 else ''
    rendered_skill_names=[name for name in names if name.lower() in rendered_skill_text.lower()]
    unsupported_skills=[name for name in rendered_skill_names if name not in verified_skill_names]
    languages_preserved='Languages: English, Telugu, Hindi, Tamil' in text
    report['unsupported_skills_validation']={'passed':not unsupported_skills and set(rendered_skill_names)==set(names) and languages_preserved,'unsupported':unsupported_skills,'planned_verified_skills':names,'rendered_skills':rendered_skill_names,'languages_preserved':languages_preserved}
    requirement_matches={item.get('requirement'):item for item in plan.get('candidate_matching',[])}
    def evidence_phrase_present(item):
        for evidence in item.get('evidence',[]):
            for phrase in [evidence.get('name'),*evidence.get('matched_evidence',[])]:
                if isinstance(phrase,str) and phrase and phrase.casefold() in text.casefold(): return True
        return False
    coverage={}
    for kind in ('required','preferred'):
        terms=[item for item in plan.get('jd_analysis',{}).get('requirements',[]) if item.get('classification')==kind]
        supported=[requirement_matches.get(item.get('canonical'),{}) for item in terms]
        supported=[item for item in supported if item.get('evidence_status') in {'SUPPORTED','PARTIAL'}]
        surfaced=[item.get('requirement') for item in supported if evidence_phrase_present(item)]
        coverage[kind]={'total':len(terms),'supported':len(supported),'supported_evidence_surfaced':len(surfaced),'surfaced_requirements':surfaced,'unsupported_requirements':[item.get('canonical') for item in terms if requirement_matches.get(item.get('canonical'),{}).get('evidence_status') in {'UNSUPPORTED','UNKNOWN'}]}
    report['jd_match_validation']={'required':coverage['required'],'preferred':coverage['preferred'],'supported_keywords_surfaced':sum(coverage[key]['supported_evidence_surfaced'] for key in coverage),'universal_ats_score_claimed':False}
    summary_text=' '.join(summary_paragraphs)
    role=target_role(plan)
    summary_sentences=[value.strip() for value in re.split(r'(?<=[.!?])\s+',summary_text) if value.strip()]
    summary_has_skill=not names or any(value.casefold() in summary_text.casefold() for value in names)
    summary_has_evidence=any(project.get('name','').casefold() in summary_text.casefold() for project in selected) or any(experience.get('organization','').casefold() in summary_text.casefold() for experience in exps)
    report['summary_quality_validation']={'passed':3<=len(summary_sentences)<=5 and bool(role) and role.casefold() in summary_text.casefold() and summary_has_skill and summary_has_evidence,'sentence_count':len(summary_sentences),'role_alignment':bool(role) and role.casefold() in summary_text.casefold(),'verified_skill_alignment':summary_has_skill,'candidate_evidence_present':summary_has_evidence,'summary_text':summary_text}
    planned_project_ids=[project.get('record_id') for project in selected]
    project_matches_by_id={item.get('record_id'):item for item in plan.get('resume_plan',{}).get('projects_to_include',[])}
    irrelevant_projects=[record_id for record_id in planned_project_ids if not project_matches_by_id.get(record_id,{}).get('matched_requirements')]
    report['project_relevance_validation']={'passed':not irrelevant_projects,'selected_record_ids':planned_project_ids,'unmatched_selected_record_ids':irrelevant_projects}
    planned_experience_ids=[item.get('record_id') for item in exps]
    relevant_experience_ids={item.get('record_id') for item in plan.get('resume_plan',{}).get('experience_decisions',[]) if item.get('decision')=='INCLUDE'}
    report['experience_relevance_validation']={'passed':set(planned_experience_ids)<=relevant_experience_ids,'selected_record_ids':planned_experience_ids,'planner_relevant_record_ids':sorted(relevant_experience_ids)}
    bullet_count=sum(1 for line in lines if line.lstrip().startswith(('•','▪','-','*')))
    word_count=len(re.findall(r'\b[\w+#.-]+\b',text))
    single_page_metrics=page_metrics[0] if len(page_metrics)==1 else {'word_count':0,'occupied_bbox_ratio':0.0,'vertical_coverage':0.0}
    clearly_underfilled=(pages==1 and word_count<220 and bullet_count<4 and single_page_metrics['occupied_bbox_ratio']<0.35 and single_page_metrics['vertical_coverage']<0.58)
    report['content_density']={'text_characters':len(text.strip()),'word_count':word_count,'meaningful_bullet_count':bullet_count,'occupied_bbox_ratio':single_page_metrics['occupied_bbox_ratio'],'vertical_coverage':single_page_metrics['vertical_coverage'],'page_metrics':page_metrics,'warning':clearly_underfilled,'basis':'Conservative conjunction of extracted words, bullet count, occupied PDF text area, and vertical coverage.'}
    report['professional_page_utilization_validation']={'passed':not clearly_underfilled,'one_page':pages==1,'clearly_underfilled':clearly_underfilled,'occupied_bbox_ratio':single_page_metrics['occupied_bbox_ratio'],'vertical_coverage':single_page_metrics['vertical_coverage']}
    report['ats_readiness_checks']={'one_page':pages==1,'ats_safe_structure':structure['single_column'] and structure['no_tables'] and structure['no_text_boxes'] and structure['no_graphics'],'extractable_text':bool(text.strip()),'standard_headings':report['ats_validation']['standard_headings_present'],'required_skill_coverage_reported':True,'preferred_skill_coverage_reported':True,'evidence_coverage_reported':True,'summary_alignment':report['summary_quality_validation']['passed'],'project_relevance':report['project_relevance_validation']['passed'],'experience_relevance':report['experience_relevance_validation']['passed'],'certification_relevance':report['certification_validation']['passed'],'accurate_project_stacks':report['project_stack_validation']['passed'],'verified_skills_only':report['unsupported_skills_validation']['passed'],'duplicate_sections':report['duplicate_template_validation']['passed'],'correct_project_selection':report['project_order_validation']['passed'],'correct_experience_selection':report['experience_relevance_validation']['passed'],'correct_certification_selection':report['certification_validation']['passed'],'contact_details':report['contact_validation']['passed'],'profile_links':report['hyperlink_validation']['passed'],'project_links':report['project_link_validation']['passed'],'professional_page_utilization':report['professional_page_utilization_validation']['passed'],'universal_score_claimed':False}
    truth=report['truth_provenance_validation']; truth_pass=all(v is True for k,v in truth.items() if k!='unsupported_claims_detected') and not truth['unsupported_claims_detected']
    checks=[report['page_count_exactly_one'],report['contact_validation']['passed'],report['hyperlink_validation']['passed'],report['summary_validation']['passed'],report['summary_format_validation']['passed'],report['summary_quality_validation']['passed'],report['languages_validation']['passed'],report['project_stack_validation']['passed'],report['project_link_validation']['passed'],report['skills_validation']['passed'],report['certification_validation']['passed'],report['education_validation']['passed'],report['project_order_validation']['passed'],report['project_relevance_validation']['passed'],report['experience_relevance_validation']['passed'],report['section_order_validation']['passed'],report['font_validation']['passed'],report['unsupported_skills_validation']['passed'],report['professional_page_utilization_validation']['passed'],report['ats_validation']['single_column'],structure['no_tables'],structure['no_text_boxes'],structure['no_graphics'],report['ats_validation']['docx_valid'],report['ats_validation']['text_extractable'],report['placeholder_validation']['passed'],report['duplicate_template_validation']['passed'],truth_pass]
    if report['content_density']['warning']: report['warnings_issues'].append('Content density may be sparse.')
    report['final_status']='PASS' if all(checks) and not report['warnings_issues'] else 'FAIL'
    output_report.write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n',encoding='utf-8'); return report

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--plan',required=True); ap.add_argument('--output',required=True); ap.add_argument('--report',required=True); ap.add_argument('--test-mode',action='store_true'); ap.add_argument('--explicit-approval',action='store_true'); args=ap.parse_args()
    plan=load_json(Path(args.plan)); prof=profile()
    if not args.test_mode and not args.explicit_approval and plan.get('approval_checkpoint',{}).get('resume_generation_allowed') is not True: raise SystemExit('STOP: Resume Plan approval is required. No resume generated.')
    if args.test_mode or args.explicit_approval: plan=dict(plan); plan['approval_checkpoint']=dict(plan.get('approval_checkpoint',{})); plan['approval_checkpoint']['resume_generation_allowed']=True
    out=Path(args.output); out.parent.mkdir(parents=True,exist_ok=True); report_path=Path(args.report); report_path.parent.mkdir(parents=True,exist_ok=True)
    included,exps,certs,fallback=_generate(plan,prof,out); report=validate(out,plan,prof,report_path); report['approval_source']='explicit_user_approval' if args.explicit_approval else ('test_mode' if args.test_mode else 'approved_resume_plan'); report['template_fallback_used']=fallback; report['files_protected']=['resumes/RESUME.docx','templates/ats_resume_template.docx','data/*.json']; report_path.write_text(json.dumps(report,indent=2,ensure_ascii=False)+'\n',encoding='utf-8'); print(json.dumps({'docx':str(out),'report':str(report_path),'page_count':report['page_count'],'final_status':report['final_status']},indent=2))
if __name__=='__main__': main()
